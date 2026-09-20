from contextlib import redirect_stdout
from copy import deepcopy
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import numpy as np
from xlsxwriter import Workbook

from apportion import apportion1d_general, threshold_drop
from dictionaries import (ADJUSTMENT_METHODS, DEFAULT_ELECTION_SETTINGS, DIVIDER_RULES,
                          ELECTION_LAW_PRESETS, QUOTA_RULES)
from electionHandler import ElectionHandler
from electionSystem import ElectionSystem
from excel_util import (fixed_seat_threshold_text, result_fractional_digits,
                        result_number_format, result_percentage_digits,
                        prepare_formats)
from noweb import load_json, load_votes, votes_to_excel
import noweb
from par_util import parallel_dir
from simulate import Sim_result, Simulation, SimulationSettings
from input_util import check_simul_settings, check_systems, normalize_system
from methods.max_const_votes import max_const_votes
from methods.switching_se import switching as swedish_switching
from methods.swedish_style_switching import switching as swedish_style_switching
from table_util import entropy
from vote_table import check_vote_table
from voting import Election
import web
from web import app


class CurrentApplicationTest(unittest.TestCase):
    def make_system(self, table, method, threshold=0):
        system = ElectionSystem()
        system.copy_info_from_votes(table)
        system['adjustment_method'] = method
        system['adjustment_threshold'] = threshold
        return system

    def make_swedish_system(self, table, divider='nordic-1.2'):
        system = self.make_system(table, 'max-const-votes', threshold=4)
        system['primary_divider'] = divider
        system['adj_determine_divider'] = divider
        system['adjustment_preparation_method'] = 'switching_se'
        system['adj_preparation_divider'] = divider
        system['adj_alloc_divider'] = 'sainte-lague'
        system['constituency_threshold'] = 12
        system['fixed_seat_national_threshold'] = 4
        system['fixed_seat_threshold_choice'] = 1
        return system

    def test_wsgi_import_initializes_simulation_state(self):
        self.assertIsInstance(noweb.SIMULATIONS, dict)

    def test_systems_default_to_comparison(self):
        self.assertTrue(ElectionSystem()['compare_with'])
        self.assertTrue(normalize_system({})['compare_with'])
        self.assertFalse(normalize_system({'compare_with': False})['compare_with'])
        response = app.test_client().post('/api/capabilities/', json=[])
        self.assertTrue(response.get_json()['election_system']['compare_with'])

    def test_legacy_fixed_seat_eligibility_is_normalized(self):
        system = normalize_system({
            'adjustment_threshold': 4,
            'fixed_seat_eligibility': 'national-or-constituency',
        })
        self.assertEqual(system['fixed_seat_national_threshold'], 4)
        self.assertEqual(system['fixed_seat_threshold_choice'], 1)
        self.assertNotIn('fixed_seat_eligibility', system)
        self.assertEqual(normalize_system({})['fixed_seat_national_threshold'], 0)
        self.assertEqual(normalize_system({})['fixed_seat_threshold_choice'], 1)
        with self.assertRaisesRegex(ValueError, 'Unknown fixed-seat eligibility'):
            normalize_system({'fixed_seat_eligibility': 'invalid'})

    def test_default_web_ports(self):
        with patch.dict(os.environ, {}, clear=True):
            with patch.object(
                    web.os, 'uname',
                    return_value=SimpleNamespace(nodename='pluto.cs.hi.is')):
                self.assertEqual(web.default_port(), '5000')
            with patch.object(
                    web.os, 'uname',
                    return_value=SimpleNamespace(nodename='workstation')):
                self.assertEqual(web.default_port(), '5001')
        with patch.dict(os.environ, {'FLASK_RUN_PORT': '5050'}):
            self.assertEqual(web.default_port(), '5050')

    def test_result_excel_precision_setting(self):
        self.assertEqual(result_fractional_digits(None), 3)
        self.assertEqual(result_fractional_digits({'fractional_digits': 2}), 2)
        self.assertEqual(result_percentage_digits(None), 2)
        self.assertEqual(result_percentage_digits({'percentage_digits': 4}), 4)
        self.assertEqual(result_number_format(2), '#,##0.00')
        self.assertEqual(result_number_format(2, percentage=True), '#,##0.00%')
        with self.assertRaisesRegex(ValueError, 'between 0 and 10'):
            result_fractional_digits({'fractional_digits': 11})
        with self.assertRaisesRegex(ValueError, 'between 0 and 10'):
            result_percentage_digits({'percentage_digits': -1})
        workbook = Workbook(BytesIO())
        formats = prepare_formats(workbook, {
            'fractional_digits': 1, 'percentage_digits': 4,
        })
        self.assertEqual(formats['cell'].num_format, '#,##0.0')
        self.assertEqual(formats['percentages'].num_format, '#,##0.0000%')
        self.assertEqual(formats['%'].num_format, '#,##0.0000%')
        workbook.close()

    def test_single_election_excel_percentage_precision(self):
        table = load_votes('../data/2-by-2-example.csv')
        system = self.make_system(table, 'max-const-vote-percentage')
        output = BytesIO()
        ElectionHandler(table, [system], True).to_xlsx(
            output, {'percentage_digits': 4})
        with ZipFile(output) as workbook:
            styles = workbook.read('xl/styles.xml').decode()
        self.assertIn('#,##0.0000%', styles)

    def test_single_election_excel_includes_vote_percentages_and_averages(self):
        import openpyxl

        table = load_votes('../data/2-by-2-example.csv')
        system = self.make_system(table, 'max-const-vote-percentage')
        handler = ElectionHandler(table, [system], True)
        with TemporaryDirectory() as directory:
            filename = Path(directory) / 'election.xlsx'
            handler.to_xlsx(filename, {
                'fractional_digits': 1,
                'percentage_digits': 4,
            })
            workbook = openpyxl.load_workbook(filename, data_only=True)
        worksheet = workbook.active
        heading_rows = {
            cell.value: cell.row
            for cell in worksheet['A'] if cell.value in {'Votes', 'Total seats'}
        }
        vote_row = heading_rows['Votes']
        total_seats_row = heading_rows['Total seats']

        self.assertEqual(worksheet.cell(vote_row + 1, 5).value, 'Vote percentage')
        self.assertEqual(worksheet.cell(vote_row + 5, 1).value, 'Vote percentage')
        self.assertAlmostEqual(worksheet.cell(vote_row + 2, 5).value, 3800 / 8000)
        self.assertAlmostEqual(worksheet.cell(vote_row + 5, 2).value, 4300 / 8000)
        self.assertEqual(worksheet.cell(vote_row + 5, 5).value, 1)
        self.assertEqual(worksheet.cell(vote_row + 2, 5).number_format, '#,##0.0000%')

        self.assertEqual(
            worksheet.cell(total_seats_row + 1, 5).value,
            'Average votes/seat',
        )
        self.assertAlmostEqual(
            worksheet.cell(total_seats_row + 2, 5).value,
            worksheet.cell(vote_row + 2, 4).value
            / worksheet.cell(total_seats_row + 2, 4).value,
        )
        self.assertEqual(worksheet.cell(total_seats_row + 2, 5).number_format,
                         '#,##0.0')
        workbook.close()

    def test_parallel_files_can_use_service_state_directory(self):
        with TemporaryDirectory() as directory:
            with patch.dict(os.environ, {"VOTING_STATE_DIR": directory}):
                self.assertEqual(parallel_dir(), Path(directory) / "pardir")

    def test_single_cpu_simulation_runs_in_background_thread(self):
        class FakeThread:
            def __init__(self, target, args):
                self.target = target
                self.args = args
                self.started = False

            def start(self):
                self.started = True

        noweb.create_SIMULATIONS()
        with patch.object(noweb, 'Thread', FakeThread), \
             patch.object(noweb, 'Simulation', return_value=object()), \
             patch.object(noweb.par_util, 'get_id', return_value='single-cpu'):
            simid = noweb.new_simulation(
                votes={}, systems=[],
                sim_settings={'cpu_count': 1},
            )

        simulation = noweb.SIMULATIONS[simid]
        self.assertEqual(simulation['kind'], 'threaded')
        self.assertTrue(simulation['thread'].started)

    def test_csv_upload_uses_uploaded_stream(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        upload = BytesIO(b'Example,fixed,adj,A,B\nI,1,0,10,20\n')
        data = {'file': (upload, 'votes.csv')}
        response = client.post('/api/votes/upload/', data=data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()['votes'], [[10, 20]])

    def test_csv_upload_reads_pruned_votes(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        upload = BytesIO(
            b'Example,fixed,adj,A,B,Pruned\n'
            b'I,1,0,10,20,3\n'
        )
        data = {'file': (upload, 'votes.csv')}
        response = client.post('/api/votes/upload/', data=data)
        self.assertEqual(response.status_code, 200)
        vote_table = response.get_json()
        self.assertEqual(vote_table['parties'], ['A', 'B'])
        self.assertEqual(vote_table['votes'], [[10, 20]])
        self.assertEqual(vote_table['pruned'], [3])

    def test_csv_upload_reads_party_names(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        upload = BytesIO(
            b'Example,fixed,adj,A,B\n'
            b'Party names,,,Alpha Party,Beta Party\n'
            b'I,1,0,10,20\n'
        )
        response = client.post('/api/votes/upload/',
                               data={'file': (upload, 'votes.csv')})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()['party_names'],
                         ['Alpha Party', 'Beta Party'])

    def test_csv_upload_reads_adjustment_seat_maxima(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        upload = BytesIO(
            b'Example,fixed,min_adj,max_adj,A,B\n'
            b'Max adj seats,,,5,,\n'
            b'I,1,2,4,10,20\n'
            b'II,1,1,-,30,40\n'
        )
        response = client.post('/api/votes/upload/',
                               data={'file': (upload, 'votes.csv')})
        self.assertEqual(response.status_code, 200)
        vote_table = response.get_json()
        self.assertEqual(vote_table['max_total_adj_seats'], 5)
        self.assertEqual(
            [constituency['max_adj_seats']
             for constituency in vote_table['constituencies']],
            [4, None],
        )

    def test_csv_upload_rejects_blank_adjustment_seat_maximum(self):
        app.config.update(TESTING=True)
        upload = BytesIO(
            b'Example,fixed,min_adj,max_adj,A\n'
            b'Max adj seats,,,1,\n'
            b'I,1,0,,10\n'
        )
        response = app.test_client().post(
            '/api/votes/upload/', data={'file': (upload, 'votes.csv')})
        self.assertIn(
            'must be a non-negative integer or - for unlimited',
            response.get_json()['error'],
        )

    def test_csv_upload_reads_national_votes_after_blank_row(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        upload = BytesIO(
            b'Example,fixed,adj,A,B\n'
            b'I,1,0,10,20\n'
            b'\n'
            b'National,0,1,30,40\n'
        )
        response = client.post('/api/votes/upload/',
                               data={'file': (upload, 'votes.csv')})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json()['party_vote_info']['votes'], [30, 40])

    def test_csv_upload_normalizes_blank_national_votes_to_zero(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        upload = BytesIO(
            b'Example,fixed,adj,A,B\n'
            b'I,1,0,10,20\n'
            b'\n'
            b'National,0,1,,\n'
        )
        response = client.post('/api/votes/upload/',
                               data={'file': (upload, 'votes.csv')})
        self.assertEqual(response.status_code, 200)
        info = response.get_json()['party_vote_info']
        self.assertEqual(info['votes'], [0, 0])
        self.assertEqual(info['total'], 0)

    def test_short_vote_file_returns_a_format_error(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        upload = BytesIO(b'Example,fixed\nI,1\n')
        response = client.post('/api/votes/upload/',
                               data={'file': (upload, 'votes.csv')})
        self.assertIn('seat columns', response.get_json()['error'])

    def test_vote_table_validation_rejects_invalid_numbers(self):
        table = load_votes('../data/2-by-2-example.csv')
        cases = [
            ('fractional votes', ('votes', 1.5)),
            ('non-finite votes', ('votes', float('nan'))),
            ('negative seats', ('seats', -1)),
        ]
        for description, (kind, value) in cases:
            with self.subTest(description):
                invalid = deepcopy(table)
                if kind == 'votes':
                    invalid['votes'][0][0] = value
                else:
                    invalid['constituencies'][0]['num_fixed_seats'] = value
                with self.assertRaises((TypeError, ValueError)):
                    check_vote_table(invalid)

    def test_specified_national_vote_name_must_be_nonblank(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['party_vote_info'] = {
            'name': ' ',
            'num_fixed_seats': 0,
            'num_adj_seats': 0,
            'votes': [10, 20],
            'specified': True,
            'pruned': 0,
        }
        with self.assertRaisesRegex(ValueError, 'non-blank'):
            check_vote_table(table)

    def test_vote_table_api_validates_edited_tables(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        table = load_votes('../data/2-by-2-example.csv')
        table['constituencies'][0]['num_fixed_seats'] = -1
        response = client.post('/api/votes/save/', json={'vote_table': table})
        self.assertIn('may not be negative', response.get_json()['error'])

        table = load_votes('../data/2-by-2-example.csv')
        table['votes'][0][0] = -1
        response = client.post('/api/votes/save/', json={'vote_table': table})
        self.assertIn('Votes may not be negative', response.get_json()['error'])

        table = load_votes('../data/2-by-2-example.csv')
        table['max_total_adj_seats'] = 5
        table['constituencies'][0]['max_adj_seats'] = 1
        table['constituencies'][1]['max_adj_seats'] = None
        response = client.post('/api/votes/save/', json={'vote_table': table})
        self.assertIn('Maximum adjustment seats may not be below the minimum',
                      response.get_json()['error'])

    def test_vote_table_api_download_can_be_uploaded(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        table = load_votes('../data/iceland-2021.csv')

        download = client.post('/api/votes/save/', json={'vote_table': table})
        self.assertEqual(download.status_code, 200)
        upload = client.post(
            '/api/votes/upload/',
            data={'file': (BytesIO(download.data), 'votes.xlsx')},
        )

        self.assertEqual(upload.status_code, 200)
        self.assertNotIn('error', upload.get_json())
        self.assertEqual(upload.get_json()['votes'], table['votes'])

    def test_swedish_vote_table_uses_adjustment_seat_bounds(self):
        table = load_votes('../data/sweden_2022.csv')
        self.assertEqual(table['max_total_adj_seats'], 39)
        self.assertTrue(all(
            constituency['num_adj_seats'] == 0
            and constituency['max_adj_seats'] is None
            for constituency in table['constituencies']
        ))

    def test_swedish_elections_are_available_as_presets(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        presets = client.get('/api/presets/').get_json()
        swedish = [preset for preset in presets
                   if preset['Country'] == 'Sweden']
        self.assertEqual(
            [preset['Year'] for preset in swedish],
            ['2014', '2018', '2022', '2026'],
        )
        for preset in swedish:
            response = client.post(
                '/api/presets/load/',
                json={'election_id': presets.index(preset)},
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(
                response.get_json()['name'],
                f"Sweden-Riksdagen-{preset['Year']}",
            )

    def test_election_law_presets_contain_rules_but_no_seat_counts(self):
        presets = {
            preset['value']: preset['settings']
            for preset in ELECTION_LAW_PRESETS if preset['settings']
        }
        self.assertEqual(
            list(presets),
            ['default', 'denmark', 'finland', 'iceland', 'norway',
             'sweden-2014', 'sweden-2018'],
        )
        self.assertEqual(
            [preset['text'] for preset in ELECTION_LAW_PRESETS[1:]],
            [
                'Denmark (2007–)',
                'Finland (1907–)',
                'Iceland (2003–)',
                'Norway (2005–)',
                'Sweden (1988–2014)',
                'Sweden (2018–)',
            ],
        )
        self.assertEqual(
            [preset.get('system_name', preset['text'])
             for preset in ELECTION_LAW_PRESETS[1:]],
            ['Denmark', 'Finland', 'Iceland', 'Norway',
             'Sweden (1988–2014)', 'Sweden'],
        )
        self.assertEqual(
            presets['default'],
            {**DEFAULT_ELECTION_SETTINGS,
             'constituency_seat_specification': 'refer'},
        )
        defaults = ElectionSystem()
        for key, value in DEFAULT_ELECTION_SETTINGS.items():
            self.assertEqual(defaults[key], value)
        self.assertEqual(
            (presets['finland']['primary_divider'],
             presets['finland']['constituency_threshold'],
             presets['finland']['adjustment_method'],
             presets['finland']['constituency_seat_specification']),
            ('dhondt', 0, 'adjustment-as-fixed', 'refer'),
        )
        self.assertEqual(
            (presets['iceland']['adjustment_method'],
             presets['iceland']['adjustment_threshold'],
             presets['iceland']['primary_divider']),
            ('icelandic-law', 5, 'dhondt'),
        )
        self.assertEqual(
            (presets['norway']['adjustment_method'],
             presets['norway']['adjustment_threshold'],
             presets['norway']['primary_divider'],
             presets['norway']['adj_alloc_divider']),
            ('norwegian-law', 4, 'nordic-1.4', 'sainte-lague'),
        )
        self.assertEqual(
            (presets['sweden-2014']['constituency_threshold'],
             presets['sweden-2014']['fixed_seat_national_threshold'],
             presets['sweden-2014']['adjustment_preparation_method'],
             presets['sweden-2014']['adj_preparation_divider'],
             presets['sweden-2014']['adjustment_method'],
             presets['sweden-2014']['constituency_seat_specification']),
            (12, 4, 'none', 'nordic-1.4',
             'max-const-votes', 'refer'),
        )
        self.assertEqual(
            (presets['sweden-2018']['constituency_threshold'],
             presets['sweden-2018']['fixed_seat_national_threshold'],
             presets['sweden-2018']['adjustment_preparation_method'],
             presets['sweden-2018']['adj_preparation_divider'],
             presets['sweden-2018']['adjustment_method'],
             presets['sweden-2018']['constituency_seat_specification']),
            (12, 4, 'switching_se', 'nordic-1.2',
             'max-const-votes', 'refer'),
        )
        forbidden = {
            'constituencies', 'num_fixed_seats', 'num_adj_seats',
            'max_adj_seats', 'max_total_adj_seats',
            'additional_adjustment_method', 'additional_adj_alloc_divider',
        }
        for settings in presets.values():
            self.assertTrue(forbidden.isdisjoint(settings))
            self.assertIsNotNone(settings['fixed_seat_national_threshold'])
        self.assertEqual(
            [settings['fixed_seat_threshold_choice'] for settings in presets.values()],
            [1, 0, 0, 0, 0, 1, 1],
        )

        client = app.test_client()
        response = client.post('/api/capabilities/', json={})
        self.assertEqual(
            response.get_json()['capabilities']['election_law_presets'],
            ELECTION_LAW_PRESETS,
        )

    def test_swedish_divisor_starts_at_1_2(self):
        generator = DIVIDER_RULES['nordic-1.2']()
        self.assertEqual([next(generator) for _ in range(4)], [1.2, 3, 5, 7])

    def test_entropy_uses_fixed_dhondt_and_sainte_lague_rules(self):
        table = load_votes('../data/2-by-2-example.csv')
        system = self.make_system(table, 'max-const-seat-share')
        system['adj_alloc_divider'] = 'danish'
        election = ElectionHandler(table, [system], True).elections[0]

        values = election.entropies()
        seats = election.results['all_const_seats']
        self.assertEqual(
            values['entropy_dhondt'],
            entropy(table['votes'], seats, DIVIDER_RULES['dhondt']),
        )
        self.assertEqual(
            values['entropy_sainte_lague'],
            entropy(
                table['votes'], seats, DIVIDER_RULES['sainte-lague']),
        )
        self.assertNotEqual(
            values['entropy_dhondt'], values['entropy_sainte_lague'])

        excel_result = election.get_result_excel()
        self.assertEqual(excel_result['entropy_dhondt'], values['entropy_dhondt'])
        self.assertEqual(
            excel_result['entropy_sainte_lague'],
            values['entropy_sainte_lague'])

        election.set_votes(np.asarray(table['votes']) * 2)
        election.assign_seats(True)
        self.assertNotEqual(election.entropies(), values)

        settings = SimulationSettings()
        settings.update(simulation_count=0, cpu_count=1)
        simulation = Simulation(settings, [system], table)
        simulation.run_and_collect_measures(table['votes'], None)
        simulated = simulation.election_handler.elections[0].entropies()
        for measure, value in simulated.items():
            self.assertEqual(simulation.stat[measure].mean(), [value])
        self.assertNotIn('entropy', simulation.stat)

    def test_greatest_relative_representation_measures_use_direct_formulas(self):
        table = load_votes('../data/2-by-2-example.csv')
        system = self.make_system(table, 'max-const-seat-share')
        settings = SimulationSettings()
        settings.update(simulation_count=0, cpu_count=1)
        simulation = Simulation(settings, [system], table)
        simulation.run_and_collect_measures(table['votes'], None)
        election = simulation.election_handler.elections[0]
        expected = max(
            seat / reference
            for seat_row, reference_row in zip(
                election.results['all_const_seats'],
                election.ref_seat_shares)
            for seat, reference in zip(seat_row, reference_row)
            if seat
        )
        expected_under = max(
            max(0, (reference - seat) / reference)
            for seat_row, reference_row in zip(
                election.results['all_const_seats'],
                election.ref_seat_shares)
            for seat, reference in zip(seat_row, reference_row)
            if reference
        )

        self.assertAlmostEqual(
            simulation.stat['max_overrepresentation'].mean()[0], expected)
        self.assertAlmostEqual(
            simulation.stat['max_underrepresentation'].mean()[0],
            expected_under)
        self.assertNotIn('min_seat_val', simulation.stat)

    def test_finnish_2015_matches_official_party_seat_totals(self):
        table = load_votes('../data/finland_2015.csv')
        settings = next(
            preset['settings'] for preset in ELECTION_LAW_PRESETS
            if preset['value'] == 'finland'
        )
        system = self.make_system(table, settings['adjustment_method'])
        system.update(deepcopy(settings))
        election = ElectionHandler(table, [system], True).elections[0]

        self.assertEqual(
            dict(zip(table['parties'], election.results['all_const_total'])),
            {
                'CENT': 49, 'SAML': 37, 'SAF': 38, 'SDP': 34,
                'GRÖNA': 15, 'VÄNST': 12, 'SFP': 9, 'KD': 5,
                'Piratp': 0, 'IP': 0, 'FKP': 0, 'Åland': 1,
            },
        )
        self.assertEqual(sum(election.results['all_const_total']), 200)

    def test_recent_finnish_elections_are_available_as_presets(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        presets = client.get('/api/presets/').get_json()
        finnish = [preset for preset in presets
                   if preset['Country'] == 'Finland']
        self.assertEqual(
            [preset['Year'] for preset in finnish],
            ['2015', '2019', '2023'],
        )
        for preset in finnish:
            response = client.post(
                '/api/presets/load/',
                json={'election_id': presets.index(preset)},
            )
            self.assertEqual(response.status_code, 200)

    def test_recent_finnish_elections_match_official_party_seat_totals(self):
        expected = {
            '2019': {
                'SDP': 40, 'PS': 39, 'KOK': 38, 'KESK': 31,
                'VIHR': 20, 'VAS': 16, 'RKP': 9, 'KD': 5,
                'LIIKE': 1, '05-FÅ': 1,
            },
            '2023': {
                'KOK': 48, 'PS': 46, 'SDP': 43, 'KESK': 23,
                'VIHR': 13, 'VAS': 11, 'RKP': 9, 'KD': 5,
                'LIIKE': 1, '05-FÅ': 1,
            },
        }
        settings = next(
            preset['settings'] for preset in ELECTION_LAW_PRESETS
            if preset['value'] == 'finland'
        )

        for year, official in expected.items():
            with self.subTest(year=year):
                table = load_votes(f'../data/finland_{year}.csv')
                system = self.make_system(table, settings['adjustment_method'])
                system.update(deepcopy(settings))
                election = ElectionHandler(table, [system], True).elections[0]
                totals = dict(zip(
                    table['parties'],
                    election.results['all_const_total'],
                ))
                self.assertEqual(
                    {party: seats for party, seats in totals.items() if seats},
                    official,
                )
                self.assertEqual(sum(totals.values()), 200)

    def test_swedish_2018_matches_official_seat_margins(self):
        table = load_votes('../data/sweden_2018.csv')
        election = ElectionHandler(
            table, [self.make_swedish_system(table)], True).elections[0]
        allocation = np.asarray(election.results['all_const_seats'])
        expected_additional = [
            4, 3, 2, 1, 2, 2, 0, 0, 0, 0, 2, 2, 1, 1, 3,
            2, 2, 1, 1, 2, 2, 3, 1, 1, 0, 0, 1, 0, 0,
        ]
        expected_party_totals = {
            'M': 70, 'C': 31, 'L': 20, 'KD': 22,
            'S': 100, 'V': 28, 'MP': 16, 'SD': 62,
        }
        self.assertEqual(
            (allocation.sum(axis=1) - election.desired_row_sums).tolist(),
            expected_additional,
        )
        self.assertEqual(
            {party: int(seats) for party, seats in
             zip(table['parties'], allocation.sum(axis=0)) if seats},
            expected_party_totals,
        )
        self.assertEqual(
            election.demo_tables[0]['steps'][0][2],
            'No switching required',
        )

    def test_swedish_2014_matches_official_party_totals_without_switching(self):
        table = load_votes('../data/sweden_2014.csv')
        settings = next(
            preset['settings'] for preset in ELECTION_LAW_PRESETS
            if preset['value'] == 'sweden-2014'
        )
        system = self.make_system(table, settings['adjustment_method'])
        system.update(deepcopy(settings))
        election = ElectionHandler(table, [system], True).elections[0]
        allocation = np.asarray(election.results['all_const_seats'])
        self.assertEqual(
            {party: int(seats) for party, seats in
             zip(table['parties'], allocation.sum(axis=0)) if seats},
            {
                'M': 84, 'C': 22, 'L': 19, 'KD': 16,
                'S': 113, 'V': 21, 'MP': 25, 'SD': 49,
            },
        )
        self.assertIsNone(election.preparation_stepbystep)

    def test_swedish_2022_matches_official_seat_margins(self):
        table = load_votes('../data/sweden_2022.csv')
        election = ElectionHandler(
            table, [self.make_swedish_system(table)], True).elections[0]
        allocation = np.asarray(election.results['all_const_seats'])
        expected_additional = [
            5, 3, 1, 2, 3, 2, 0, 0, 0, 0, 1, 1, 2, 0, 2,
            1, 3, 0, 1, 1, 2, 3, 0, 2, 2, 1, 0, 1, 0,
        ]
        expected_party_totals = {
            'M': 68, 'C': 24, 'L': 16, 'KD': 19,
            'S': 107, 'V': 24, 'MP': 18, 'SD': 73,
        }
        self.assertEqual(
            (allocation.sum(axis=1) - election.desired_row_sums).tolist(),
            expected_additional,
        )
        self.assertEqual(
            {party: int(seats) for party, seats in
            zip(table['parties'], allocation.sum(axis=0)) if seats},
            expected_party_totals,
        )
        self.assertEqual(len(election.demo_tables), 2)
        self.assertEqual(
            election.demo_tables[0]['headers'],
            [
                'Step', 'Constituency', 'From', 'To',
                'Returned quotient', 'Recipient quotient',
            ],
        )
        self.assertEqual(
            election.demo_tables[0]['steps'][0][2],
            'No switching required',
        )
        self.assertEqual(len(election.demo_tables[1]['steps']), 39)
        web_result = election.get_result_web()
        display = web_result['display_results']
        self.assertEqual(
            display[-1][table['parties'].index('M')],
            '68 (1)',
        )
        self.assertEqual(display[-1][-1], '349 (39)')
        self.assertFalse(web_result['switching_affected'])
        unchanged = np.argwhere(
            (np.asarray(election.results['all_const_seats']) > 0)
            & (election.switching_seat_changes == 0)
            & (election.adjustment_seat_allocations == 0)
        )[0]
        c, p = map(int, unchanged)
        self.assertEqual(
            display[c][p], str(election.results['all_const_seats'][c][p]))

    def test_swedish_switching_marks_only_affected_lists(self):
        self.assertEqual(Election.display_swedish_seats(0, 0, True), '0*')
        self.assertEqual(Election.display_swedish_seats(1, 1, True), '1 (1)*')
        table = load_votes('../data/sweden_2022.csv')
        election = ElectionHandler(
            table, [self.make_swedish_system(table)], True).elections[0]
        election.switching_seat_changes[0, 0] = -1
        election.switching_seat_changes[0, 1] = 1
        web_result = election.get_result_web()
        display = web_result['display_results']
        self.assertTrue(web_result['switching_affected'])
        for party_index in (0, 1):
            expected = election.display_seats(
                election.results['all'][0][party_index],
                election.adjustment_seat_allocations[0][party_index])
            self.assertEqual(display[0][party_index], f'{expected}*')
        self.assertFalse(display[0][-1].endswith('*'))
        self.assertFalse(display[-1][0].endswith('*'))
        self.assertFalse(display[-1][-1].endswith('*'))

    def test_swedish_2026_matches_official_party_totals(self):
        table = load_votes('../data/sweden_2026.csv')
        self.assertGreater(len(table['parties']), 8)
        self.assertEqual(sum(table['pruned']), 0)
        self.assertEqual(np.asarray(table['votes'])[:, 8:].sum(), 107_199)
        election = ElectionHandler(
            table, [self.make_swedish_system(table)], True).elections[0]
        allocation = np.asarray(election.results['all_const_seats'])
        self.assertEqual(
            {party: int(seats) for party, seats in
             zip(table['parties'], allocation.sum(axis=0)) if seats},
            {
                'M': 70, 'C': 25, 'L': 19, 'KD': 22,
                'S': 99, 'V': 30, 'MP': 22, 'SD': 62,
            },
        )

    def test_swedish_system_runs_through_simulation_measures(self):
        table = load_votes('../data/sweden_2022.csv')
        settings = SimulationSettings()
        settings['simulation_count'] = 1
        settings['cpu_count'] = 1
        with redirect_stdout(StringIO()):
            simulation = Simulation(
                settings, [self.make_swedish_system(table)], table)
            simulation.run_and_collect_measures(table['votes'], None)
        self.assertEqual(
            sum(simulation.election_handler.elections[0]
                .results['all_const_total']),
            349,
        )

    def test_disabled_thresholds_apply_to_counterfactual_simulation_measures(self):
        table = load_votes('../data/2-by-2-example.csv')
        threshold_system = self.make_system(
            table, 'max-const-seat-share', threshold=50)
        zero_system = deepcopy(threshold_system)
        zero_system['adjustment_threshold'] = 0

        disabled_settings = SimulationSettings()
        disabled_settings.update({
            'simulation_count': 0,
            'cpu_count': 1,
            'use_thresholds': False,
        })
        zero_settings = deepcopy(disabled_settings)
        zero_settings['use_thresholds'] = True

        disabled = Simulation(disabled_settings, [threshold_system], table)
        explicit_zero = Simulation(zero_settings, [zero_system], table)
        for simulation in (disabled, explicit_zero):
            simulation.run_and_collect_measures(table['votes'], None)

        self.assertEqual(disabled.stat['dev_all_adj_const'].mean(), [0.0])
        self.assertEqual(disabled.stat['dev_all_adj_tot'].mean(), [0.0])
        self.assertEqual(
            disabled.stat['dev_all_adj_const'].mean(),
            explicit_zero.stat['dev_all_adj_const'].mean(),
        )
        self.assertEqual(
            disabled.stat['dev_all_adj_tot'].mean(),
            explicit_zero.stat['dev_all_adj_tot'].mean(),
        )

    def test_random_seed_reproduces_vote_generation_across_worker_ranges(self):
        table = load_votes('../data/2-by-2-example.csv')
        system = self.make_system(table, 'max-const-seat-share')
        settings = SimulationSettings()
        settings.update(random_seed=24680, simulation_count=3, cpu_count=1)

        for distribution in ('log-normal', 'uniform', 'gamma', 'beta'):
            with self.subTest(distribution=distribution):
                settings['gen_method'] = distribution
                full = Simulation(settings, [system], table)
                first_worker = Simulation(settings, [self.make_system(
                    table, 'max-const-seat-share')], table, start_iteration=0)
                second_worker = Simulation(settings, [self.make_system(
                    table, 'max-const-seat-share')], table, start_iteration=2)

                expected = [full.generate_simulated_votes(i) for i in range(3)]
                actual = [first_worker.generate_simulated_votes(i) for i in range(2)]
                actual.append(second_worker.generate_simulated_votes(2))
                self.assertEqual(expected, actual)

    def test_comparison_measures_are_combined_once_across_workers(self):
        table = load_votes('../data/2-by-2-example.csv')
        systems = []
        for name, divider in (("D'Hondt", 'dhondt'),
                              ('Sainte-Laguë', 'sainte-lague')):
            system = self.make_system(table, 'max-const-seat-share')
            system['name'] = name
            system['primary_divider'] = divider
            system['adj_determine_divider'] = divider
            system['adj_alloc_divider'] = divider
            systems.append(system)

        settings = SimulationSettings()
        settings.update(random_seed=12345, simulation_count=9, cpu_count=2)

        def run(count, start_iteration):
            worker_settings = deepcopy(settings)
            worker_settings['simulation_count'] = count
            simulation = Simulation(
                worker_settings, deepcopy(systems), deepcopy(table),
                start_iteration=start_iteration)
            simulation.simulate(tasknr=1)
            return Sim_result(simulation.attributes())

        uninterrupted = run(9, 0)
        combined = run(4, 0)
        combined.combine(run(5, 4))

        comparison_measures = [
            measure for measure in uninterrupted.MEASURES
            if measure.startswith('cmp_')
        ]
        self.assertTrue(comparison_measures)
        for measure in comparison_measures:
            with self.subTest(measure=measure):
                self.assertEqual(combined.stat[measure].n, 9)
                np.testing.assert_allclose(
                    combined.stat[measure].numpy_mean(),
                    uninterrupted.stat[measure].numpy_mean(),
                    atol=1e-12)
                np.testing.assert_allclose(
                    combined.stat[measure].numpy_std(),
                    uninterrupted.stat[measure].numpy_std(),
                    atol=1e-12)

        for simulation in (uninterrupted, combined):
            values = simulation.stat['sum_abs'].numpy_mean()
            self.assertEqual(len(values), 3)
            self.assertAlmostEqual(values[2], values[0] - values[1])

        uninterrupted.analysis()
        web_result = uninterrupted.get_result_web(False)
        paired = web_result['paired_data']['sum_abs']
        self.assertAlmostEqual(
            paired['avg'],
            web_result['data'][0]['measures']['sum_abs']['avg']
            - web_result['data'][1]['measures']['sum_abs']['avg'],
        )
        displayed = web_result['vuedata']['toLists'][0]['avg']
        self.assertEqual(len(displayed), 3)
        self.assertAlmostEqual(displayed[2]['value'], paired['avg'])
        self.assertIn(
            "D'Hondt minus Sainte-Laguë",
            web_result['vuedata']['difference_tooltip'],
        )

    def test_random_seed_validation(self):
        for seed in ('', '-'):
            settings = SimulationSettings()
            settings['random_seed'] = seed
            self.assertIsNone(check_simul_settings(settings)['random_seed'])
        settings = SimulationSettings()
        settings['random_seed'] = 123
        self.assertEqual(check_simul_settings(settings)['random_seed'], 123)
        settings = SimulationSettings()
        settings['random_seed'] = 2**31
        with self.assertRaisesRegex(ValueError, 'Random seed'):
            check_simul_settings(settings)

    def test_simulation_endpoint_accepts_cleared_random_seed(self):
        table = load_votes('../data/2-by-2-example.csv')
        system = self.make_system(table, 'max-const-seat-share')
        settings = SimulationSettings()
        settings['random_seed'] = '-'
        with patch.object(web, 'new_simulation', return_value='test-id') as start:
            response = app.test_client().post('/api/simulate/', json={
                'vote_table': table,
                'systems': [system],
                'sim_settings': settings,
            })
            self.assertEqual(response.get_json(), {'started': True, 'simid': 'test-id'})
            self.assertIsNone(start.call_args.args[2]['random_seed'])

            settings['random_seed'] = ''
            response = app.test_client().post('/api/simulate/', json={
                'vote_table': table,
                'systems': [system],
                'sim_settings': settings,
            })
            self.assertEqual(response.get_json(), {'started': True, 'simid': 'test-id'})
            self.assertIsNone(start.call_args.args[2]['random_seed'])

            settings['random_seed'] = 'not an integer'
            response = app.test_client().post('/api/simulate/', json={
                'vote_table': table,
                'systems': [system],
                'sim_settings': settings,
            })
            self.assertIn('Random seed must be an integer', response.get_json()['error'])
            self.assertEqual(start.call_count, 2)

    def test_bulk_vote_generation_with_national_votes(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['party_vote_info'] = {
            'name': 'National votes',
            'num_fixed_seats': 0,
            'num_adj_seats': 0,
            'votes': [4200, 3800],
            'specified': True,
            'pruned': 0,
        }
        for distribution in ('uniform', 'gamma', 'beta'):
            with self.subTest(distribution=distribution):
                settings = SimulationSettings()
                settings.update(random_seed=42, cpu_count=1, gen_method=distribution)
                system = self.make_system(table, 'max-const-seat-share')
                simulation = Simulation(settings, [system], table)
                votes, national = simulation.generate_simulated_votes(0)
                self.assertEqual(len(national), 2)
                self.assertTrue(all(v > 0 for v in national))
                self.assertEqual(
                    (votes, national), simulation.generate_simulated_votes(0))
                simulation.run_and_collect_measures(votes, national, 0)

    def test_swedish_switching_returns_excess_seat(self):
        allocation, steps = swedish_switching(
            [[1, 1], [24, 1]],
            [3, 3],
            [3, 3],
            np.array([[1, 2], [3, 0]], dtype=int),
            DIVIDER_RULES['nordic-1.2'],
            nat_votes=np.array([25, 2]),
            total_seats=6,
        )
        np.testing.assert_array_equal(allocation, [[0, 3], [3, 0]])
        np.testing.assert_array_equal(steps['party_totals'], [3, 3])
        self.assertEqual(len(steps['data']['switches']), 1)

    def test_swedish_style_switching_protects_fixed_seats(self):
        prior = np.array([[0, 0], [1, 0]], dtype=int)
        allocation, steps = swedish_style_switching(
            [[10, 0], [100, 0]], [1, 2], [2, 1], prior,
            DIVIDER_RULES['sainte-lague'],
        )
        np.testing.assert_array_equal(allocation, [[0, 1], [2, 0]])
        self.assertTrue(np.all(allocation >= prior))
        self.assertEqual(len(steps['data']['switches']), 1)
        self.assertEqual(steps['data']['switches'][0]['constituency'], 0)

    def test_swedish_style_switching_rejects_fixed_seat_excess(self):
        with self.assertRaisesRegex(ValueError, 'No removable excess seat'):
            swedish_style_switching(
                [[10, 1]], [2], [1, 1], [[2, 0]],
                DIVIDER_RULES['sainte-lague'],
            )

    def test_max_const_votes_respects_constituency_capacity(self):
        allocation, _ = max_const_votes(
            [[100, 10], [90, 9]],
            [1, 3],
            [2, 2],
            np.array([[1, 0], [1, 0]]),
            DIVIDER_RULES['sainte-lague'],
            num_adjustment_seats=2,
            min_adj_seats=[0, 2],
            max_adj_seats=[0, 2],
        )
        np.testing.assert_array_equal(allocation, [[1, 0], [1, 2]])

    def test_max_const_votes_respects_constituency_minimums(self):
        allocation, _ = max_const_votes(
            [[1, 1], [100, 100]],
            [2, 0],
            [2, 2],
            np.zeros((2, 2), dtype=int),
            DIVIDER_RULES['sainte-lague'],
            num_adjustment_seats=4,
            min_adj_seats=[2, 0],
            max_adj_seats=[None, None],
        )
        np.testing.assert_array_equal(allocation.sum(axis=1), [2, 2])
        np.testing.assert_array_equal(allocation.sum(axis=0), [2, 2])

    def test_max_const_votes_handles_flexible_adjustment_seat_bounds(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['constituencies'][0]['max_adj_seats'] = 3
        table['constituencies'][1]['max_adj_seats'] = 4
        table['max_total_adj_seats'] = 7
        system = self.make_system(table, 'max-const-votes')

        election = ElectionHandler(table, [system], True).elections[0]
        adjustment = (
            np.asarray(election.results['all_const_seats'])
            - np.asarray(election.results['fixed_const_seats']))

        np.testing.assert_array_equal(adjustment.sum(axis=1), [3, 4])
        self.assertEqual(int(adjustment.sum()), 7)

    def test_ordinary_method_rejects_flexible_adjustment_seat_bounds(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['constituencies'][0]['max_adj_seats'] = 3
        table['constituencies'][1]['max_adj_seats'] = None
        table['max_total_adj_seats'] = 6
        for method in ADJUSTMENT_METHODS:
            if method in ('max-const-votes', 'max-const-vote-percentage'):
                continue
            with self.subTest(method=method):
                system = self.make_system(table, method)
                with self.assertRaisesRegex(
                        ValueError, 'does not support constituency ranges'):
                    ElectionHandler(table, [system], True)

    def test_ordinary_method_accepts_explicit_exact_adjustment_seats(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['constituencies'][0]['max_adj_seats'] = 2
        table['constituencies'][1]['max_adj_seats'] = 3
        table['max_total_adj_seats'] = 5
        system = self.make_system(table, 'max-const-seat-share')

        election = ElectionHandler(table, [system], True).elections[0]

        np.testing.assert_array_equal(election.final_row_sums, [12, 13])

    def test_generated_party_names_are_included(self):
        cases = [
            ('../data/norway_2017.csv', 'AP', 'Arbeiderpartiet'),
            ('../data/norway_2021.csv', 'AP', 'Arbeiderpartiet'),
            ('../data/norway_2025.csv', 'AP', 'Arbeiderpartiet'),
            ('../data/sweden_2018.csv', 'M', 'Moderaterna'),
        ]
        for filename, abbreviation, name in cases:
            with self.subTest(filename=filename):
                table = load_votes(filename)
                party_index = table['parties'].index(abbreviation)
                self.assertEqual(table['party_names'][party_index], name)

    def test_blank_party_names_are_omitted(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['party_names'] = ['', '  ']
        self.assertNotIn('party_names', check_vote_table(table))

    def test_party_names_must_match_the_party_list(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['party_names'] = ['Only one']
        with self.assertRaisesRegex(ValueError, 'Party names'):
            check_vote_table(table)

    def test_single_election_xlsx_includes_party_names(self):
        import openpyxl

        table = load_votes('../data/iceland-2021.csv')
        system = self.make_system(table, 'max-const-seat-share')
        handler = ElectionHandler(table, [system], use_thresholds=True)
        with TemporaryDirectory() as directory:
            filename = Path(directory) / 'election.xlsx'
            handler.to_xlsx(filename)
            workbook = openpyxl.load_workbook(filename, read_only=True)
        worksheet = workbook['Party names']
        self.assertEqual(
            [worksheet['A1'].value, worksheet['B1'].value],
            ['Abbreviation', 'Name'],
        )
        self.assertEqual(
            [worksheet['A2'].value, worksheet['B2'].value],
            ['B', 'Framsóknarflokkur'],
        )

    def test_vote_table_xlsx_round_trip_preserves_party_names(self):
        table = load_votes('../data/iceland-2021.csv')
        with TemporaryDirectory() as directory:
            filename = Path(directory) / 'votes.xlsx'
            votes_to_excel(table, filename)
            loaded = load_votes(filename)
        self.assertEqual(loaded['party_names'], table['party_names'])

    def test_vote_table_xlsx_round_trip_preserves_adjustment_seat_maxima(self):
        import openpyxl

        table = load_votes('../data/2-by-2-example.csv')
        table['max_total_adj_seats'] = 6
        table['constituencies'][0]['max_adj_seats'] = 2
        table['constituencies'][1]['max_adj_seats'] = None
        with TemporaryDirectory() as directory:
            filename = Path(directory) / 'votes.xlsx'
            votes_to_excel(table, filename)
            workbook = openpyxl.load_workbook(filename, read_only=True)
            self.assertEqual(workbook.active['D4'].value, '-')
            workbook.close()
            loaded = load_votes(filename)
        self.assertEqual(loaded['max_total_adj_seats'], 6)
        self.assertEqual(
            [constituency['max_adj_seats']
             for constituency in loaded['constituencies']],
            [2, None],
        )

    def test_excel_round_trip_preserves_pruned_votes(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['pruned'] = [3, 4]
        table['party_vote_info'] = {
            'name': 'National',
            'num_fixed_seats': 0,
            'num_adj_seats': 0,
            'votes': [12, 21],
            'specified': True,
            'total': 37,
            'pruned': 4,
        }
        with TemporaryDirectory() as directory:
            filename = Path(directory) / 'votes.xlsx'
            votes_to_excel(table, filename)
            loaded = load_votes(filename)
        self.assertEqual(loaded['pruned'], [3, 4])
        self.assertEqual(loaded['party_vote_info']['pruned'], 4)

    def test_exact_percentage_threshold_qualifies(self):
        allocation, _, _ = apportion1d_general(
            v_votes=[400, 600],
            num_total_seats=10,
            prior_allocations=[],
            rule=DIVIDER_RULES['dhondt'],
            threshold_percent=40,
        )
        self.assertEqual(allocation.tolist(), [4, 6])

    def test_percentage_threshold_can_use_complete_vote_total(self):
        threshold = [0, 4, 0, None]
        self.assertEqual(threshold_drop([39, 936], threshold), [39, 936])
        self.assertEqual(
            threshold_drop([39, 936], threshold, threshold_total=1000),
            [0, 936],
        )

    def test_quota_includes_pruned_but_not_threshold_excluded_votes(self):
        cases = [
            ('hare', 36),
            ('droop', 70 - 170 / 6),
        ]
        for rule_name, second_value in cases:
            with self.subTest(rule=rule_name):
                _, seat_factory, _ = apportion1d_general(
                    v_votes=[70, 30],
                    num_total_seats=5,
                    prior_allocations=[],
                    rule=QUOTA_RULES[rule_name],
                    type_of_rule='Quota',
                    threshold_percent=30,
                    threshold_total=200,
                )
                seats = seat_factory()
                self.assertEqual(next(seats)['active_votes'], 70)
                self.assertAlmostEqual(next(seats)['active_votes'], second_value)

    def test_constituency_threshold_includes_pruned_votes(self):
        table = {
            'name': 'Threshold example',
            'parties': ['A', 'B'],
            'votes': [[39, 936]],
            'pruned': [25],
            'constituencies': [{
                'name': 'I',
                'num_fixed_seats': 100,
                'num_adj_seats': 0,
            }],
            'party_vote_info': {
                'name': '-',
                'num_fixed_seats': 0,
                'num_adj_seats': 0,
                'votes': [],
                'specified': False,
                'total': 0,
                'pruned': 0,
            },
        }
        system = self.make_system(table, 'max-const-seat-share', threshold=0)
        system['constituency_threshold'] = 4
        system['fixed_seat_threshold_choice'] = 0
        handler = ElectionHandler(table, [system], use_thresholds=True)
        allocation = handler.elections[0].results['fixed_const_seats'][0]
        self.assertEqual(allocation, [0, 100])

    def test_fixed_seat_threshold_combination(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['constituencies'] = table['constituencies'][:1]
        table['constituencies'][0].update(num_fixed_seats=2, num_adj_seats=0)
        table['votes'] = [[40, 60]]
        table['pruned'] = [0]
        system = self.make_system(table, 'max-const-seat-share', threshold=90)
        system['constituency_threshold'] = 50
        for national_threshold, choice, expected in (
                (0, 0, [0, 2]), (40, 0, [0, 2]),
                (40, 1, [1, 1]), (0, 1, [1, 1])):
            with self.subTest(national_threshold=national_threshold, choice=choice):
                system['fixed_seat_national_threshold'] = national_threshold
                system['fixed_seat_threshold_choice'] = choice
                allocation = ElectionHandler(
                    table, [system], use_thresholds=True).elections[0]
                self.assertEqual(allocation.results['fixed_const_seats'][0], expected)
        self.assertEqual(fixed_seat_threshold_text(system), '0% national or 50% local')
        system['fixed_seat_threshold_choice'] = 0
        self.assertEqual(fixed_seat_threshold_text(system), '0% national and 50% local')
        table['pruned'] = [20]
        for national_threshold, expected in ((40, [0, 2]), (30, [1, 1])):
            with self.subTest(pruned_national_threshold=national_threshold):
                system['fixed_seat_national_threshold'] = national_threshold
                system['fixed_seat_threshold_choice'] = 1
                allocation = ElectionHandler(
                    table, [system], use_thresholds=True).elections[0]
                self.assertEqual(allocation.results['fixed_const_seats'][0], expected)

    def test_fixed_seat_national_threshold_validation(self):
        table = load_votes('../data/2-by-2-example.csv')
        system = self.make_system(table, 'max-const-seat-share')
        for value in (0, 100):
            with self.subTest(value=value):
                system['fixed_seat_national_threshold'] = value
                check_systems([system])
        for value in (None, '', -1, 101, 'invalid', True):
            with self.subTest(value=value):
                system['fixed_seat_national_threshold'] = value
                with self.assertRaisesRegex(ValueError, 'National threshold'):
                    check_systems([system])
        system['fixed_seat_national_threshold'] = 0
        for value in (None, '', -1, 101, 'invalid', True):
            with self.subTest(adjustment_threshold=value):
                system['adjustment_threshold'] = value
                with self.assertRaisesRegex(ValueError, 'National threshold'):
                    check_systems([system])
        system['adjustment_threshold'] = 0
        for value in (None, '', 2, True):
            with self.subTest(fixed_seat_threshold_choice=value):
                system['fixed_seat_threshold_choice'] = value
                with self.assertRaisesRegex(ValueError, 'threshold combination'):
                    check_systems([system])

    def test_zero_list_votes_are_preserved_in_election_calculations(self):
        table = {
            'name': 'Zero-list example',
            'parties': ['A', 'B'],
            'votes': [[0, 100], [100, 0]],
            'constituencies': [
                {'name': 'I', 'num_fixed_seats': 1, 'num_adj_seats': 0},
                {'name': 'II', 'num_fixed_seats': 1, 'num_adj_seats': 0},
            ],
            'party_vote_info': {
                'name': '-',
                'num_fixed_seats': 0,
                'num_adj_seats': 0,
                'votes': [],
                'specified': False,
                'total': 0,
                'pruned': 0,
            },
        }
        system = self.make_system(table, 'max-const-seat-share')
        election = ElectionHandler(
            table, [system], use_thresholds=True).elections[0]

        np.testing.assert_array_equal(
            election.votes, [[0, 100], [100, 0]])
        np.testing.assert_array_equal(election.votesums, [100, 100])
        np.testing.assert_array_equal(
            election.const_threshold_totals, [100, 100])

    def test_common_allocation_does_not_add_forced_steps(self):
        table = {
            'name': 'Adjustment example',
            'parties': ['A', 'B'],
            'votes': [[100, 0], [0, 100]],
            'constituencies': [
                {'name': 'I', 'num_fixed_seats': 0, 'num_adj_seats': 1},
                {'name': 'II', 'num_fixed_seats': 0, 'num_adj_seats': 1},
            ],
            'party_vote_info': {
                'name': '-',
                'num_fixed_seats': 0,
                'num_adj_seats': 0,
                'votes': [],
                'specified': False,
                'total': 0,
                'pruned': 0,
            },
        }
        system = self.make_system(table, 'max-const-seat-share')
        election = ElectionHandler(
            table, [system], use_thresholds=True).elections[0]

        criteria_column = election.demo_tables[0]['headers'].index('Criteria')
        self.assertEqual(
            [row[criteria_column] for row in election.demo_tables[0]['steps']],
            ['Max over all lists', 'Max over all lists'],
        )

    def test_common_allocation_uses_logical_votes_without_changing_totals(self):
        table = {
            'name': 'Logical-vote example',
            'parties': ['A', 'B'],
            'votes': [[100, 0], [100, 0]],
            'constituencies': [
                {'name': 'I', 'num_fixed_seats': 0, 'num_adj_seats': 1},
                {'name': 'II', 'num_fixed_seats': 0, 'num_adj_seats': 1},
            ],
            'party_vote_basis': 'party_vote_info',
            'party_vote_info': {
                'name': 'National',
                'num_fixed_seats': 0,
                'num_adj_seats': 0,
                'votes': [100, 100],
                'specified': True,
                'total': 200,
                'pruned': 0,
            },
        }
        system = self.make_system(table, 'max-const-seat-share')
        election = ElectionHandler(
            table, [system], use_thresholds=True).elections[0]

        np.testing.assert_array_equal(election.votes, table['votes'])
        np.testing.assert_array_equal(election.votesums, [200, 0])
        np.testing.assert_array_equal(
            election.results['all_const_seats'], [[1, 0], [0, 1]])

    def test_parties_can_be_required_to_stand_in_every_constituency(self):
        table = {
            'name': 'Norwegian standing example',
            'parties': ['A', 'B'],
            'votes': [[60, 100], [60, 0]],
            'constituencies': [
                {'name': 'I', 'num_fixed_seats': 1, 'num_adj_seats': 1},
                {'name': 'II', 'num_fixed_seats': 1, 'num_adj_seats': 1},
            ],
            'party_vote_info': {
                'name': '-',
                'num_fixed_seats': 0,
                'num_adj_seats': 0,
                'votes': [],
                'specified': False,
                'total': 0,
                'pruned': 0,
            },
        }
        system = self.make_system(table, 'max-const-votes', threshold=4)
        system['primary_divider'] = 'dhondt'
        system['adj_determine_divider'] = 'dhondt'
        system['adj_alloc_divider'] = 'dhondt'
        unrestricted = ElectionHandler(table, [system], use_thresholds=True).elections[0]
        np.testing.assert_array_equal(unrestricted.desired_col_sums, [2, 2])

        system['require_votes_in_all_constituencies'] = True
        handler = ElectionHandler(table, [system], use_thresholds=True)
        election = handler.elections[0]

        np.testing.assert_array_equal(election.party_stands_everywhere, [True, False])
        np.testing.assert_array_equal(election.desired_col_sums, [3, 1])
        np.testing.assert_array_equal(election.results['all_const_total'], [3, 1])

        # Standing comes from the source table and does not change when a
        # simulated vote gives B positive votes in the second constituency.
        handler.run_elections(True, [[60, 100], [60, 50]])
        np.testing.assert_array_equal(election.desired_col_sums, [3, 1])

    def test_norwegian_preset_requires_standing_everywhere(self):
        presets = {preset['value']: preset['settings']
                   for preset in ELECTION_LAW_PRESETS}
        self.assertTrue(presets['norway']['require_votes_in_all_constituencies'])

    def test_national_threshold_totals_include_matching_pruned_votes(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['pruned'] = [100, 200]
        table['party_vote_info'] = {
            'name': 'National',
            'num_fixed_seats': 0,
            'num_adj_seats': 0,
            'votes': [4000, 4000],
            'specified': True,
            'total': 8400,
            'pruned': 400,
        }
        expected_totals = {
            'totals': 8300,
            'party_vote_info': 8400,
            'average': 8350,
        }
        for basis, expected in expected_totals.items():
            with self.subTest(basis=basis):
                table['party_vote_basis'] = basis
                system = self.make_system(table, 'max-const-seat-share')
                handler = ElectionHandler(
                    table, [system], use_thresholds=True)
                self.assertEqual(
                    handler.elections[0].nat_threshold_total, expected)

    def test_national_threshold_excludes_party_below_complete_vote_share(self):
        table = {
            'name': 'National threshold example',
            'parties': ['A', 'B'],
            'votes': [[500, 500]],
            'pruned': [0],
            'constituencies': [{
                'name': 'I',
                'num_fixed_seats': 0,
                'num_adj_seats': 100,
            }],
            'party_vote_info': {
                'name': 'National',
                'num_fixed_seats': 0,
                'num_adj_seats': 0,
                'votes': [39, 936],
                'specified': True,
                'total': 1000,
                'pruned': 25,
            },
        }
        table['party_vote_basis'] = 'party_vote_info'
        system = self.make_system(table, 'max-const-seat-share', threshold=4)
        handler = ElectionHandler(table, [system], use_thresholds=True)
        self.assertEqual(
            handler.elections[0].desired_col_sums.tolist(), [0, 100])

    def test_simulated_votes_keep_recorded_pruned_threshold_total(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['pruned'] = [100, 200]
        system = self.make_system(table, 'max-const-seat-share')
        handler = ElectionHandler(table, [system], use_thresholds=True)
        generated = [[1000, 2000], [3000, 4000]]
        handler.run_elections(True, votes=generated)
        self.assertEqual(
            handler.elections[0].const_threshold_totals.tolist(),
            [3100, 7200],
        )

    def test_run_elections_accepts_numpy_vote_matrix(self):
        table = load_votes('../data/2-by-2-example.csv')
        system = self.make_system(table, 'max-const-seat-share')
        handler = ElectionHandler(table, [system], use_thresholds=True)

        handler.run_elections(True, votes=np.asarray(table['votes']))

        self.assertEqual(
            sum(handler.elections[0].results['all_const_total']), 25)

    def test_old_vote_tables_default_to_no_pruned_votes(self):
        table = load_votes('../data/2-by-2-example.csv')
        table.pop('pruned')
        table['party_vote_info'].pop('pruned')
        checked = check_vote_table(table)
        self.assertEqual(checked['pruned'], [0, 0])
        self.assertEqual(checked['party_vote_info']['pruned'], 0)
        self.assertEqual(checked['party_vote_basis'], 'totals')

    def test_party_vote_basis_defaults_to_totals_without_party_votes(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['party_vote_basis'] = 'party_vote_info'
        checked = check_vote_table(table)
        self.assertEqual(checked['party_vote_basis'], 'totals')

    def test_average_national_vote_basis(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['party_vote_info'] = {
            'name': 'National',
            'num_fixed_seats': 0,
            'num_adj_seats': 0,
            'votes': [4000, 4000],
            'specified': True,
            'total': 8000,
        }
        table['party_vote_basis'] = 'average'
        system = self.make_system(table, 'max-const-seat-share')
        handler = ElectionHandler(table, [system], use_thresholds=True)
        self.assertEqual(handler.elections[0].nat_votes.tolist(), [4150, 3850])

    def test_party_vote_basis_applies_to_all_systems(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['party_vote_info'] = {
            'name': 'National',
            'num_fixed_seats': 0,
            'num_adj_seats': 0,
            'votes': [4200, 3800],
            'specified': True,
            'total': 8000,
            'pruned': 0,
        }
        table['party_vote_basis'] = 'party_vote_info'
        first = self.make_system(table, 'max-const-seat-share')
        second = self.make_system(table, 'max-const-seat-share')
        second['seat_spec_options']['party'] = 'average'
        handler = ElectionHandler(table, [first, second], use_thresholds=True)
        for election in handler.elections:
            self.assertEqual(election.nat_votes.tolist(), [4200, 3800])
            self.assertEqual(
                election.system['seat_spec_options']['party'],
                'party_vote_info',
            )

    def test_old_json_uses_first_system_party_vote_basis(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['party_vote_info'] = {
            'name': 'National',
            'num_fixed_seats': 0,
            'num_adj_seats': 0,
            'votes': [4000, 4000],
            'specified': True,
            'total': 8000,
            'pruned': 0,
        }
        table.pop('party_vote_basis')
        system = self.make_system(table, 'max-const-seat-share')
        system['seat_spec_options']['party'] = 'average'
        contents = {
            'vote_table': table,
            'systems': [system],
            'sim_settings': SimulationSettings(),
        }
        with TemporaryDirectory() as directory:
            filename = Path(directory) / 'download-all.json'
            filename.write_text(json.dumps(contents), encoding='utf-8')
            loaded = load_json(filename)
        self.assertEqual(loaded['vote_table']['party_vote_basis'], 'average')

    def test_settings_download_preserves_adjustment_preparation(self):
        table = load_votes('../data/sweden_2018.csv')
        system = self.make_swedish_system(table)
        system['compare_with'] = False
        system['nat_seats'] = {
            'specified': False,
            'num_fixed_seats': 0,
            'num_adj_seats': 0,
        }
        app.config.update(TESTING=True)
        response = app.test_client().post('/api/settings/save/', json={
            'systems': [system],
            'sim_settings': SimulationSettings(),
        })
        saved = json.loads(response.data)
        saved_system = saved['e_settings'][0]
        self.assertEqual(
            saved_system['adjustment_preparation_method'], 'switching_se')
        self.assertEqual(
            saved_system['adjustment_preparation_rule'], 'nordic-1.2')
        self.assertEqual(
            saved_system['fixed_seat_national_threshold'], 4)
        self.assertEqual(saved_system['danish_special_rules'], False)
        self.assertEqual(
            saved_system['adjustment_threshold_seats'],
            system['adjustment_threshold_seats'])
        self.assertEqual(
            saved_system['adj_threshold_choice'],
            system['adj_threshold_choice'])
        self.assertNotIn('additional_adjustment_method', saved_system)
        self.assertNotIn('additional_adjustment_allocation_rule', saved_system)

    def test_fractional_reference_ignores_thresholds_and_divider_rules(self):
        table = {
            'name': 'Fractional reference example',
            'parties': ['A', 'B'],
            'votes': [[60, 40], [60, 40]],
            'pruned': [0, 0],
            'constituencies': [
                {'name': 'I', 'num_fixed_seats': 0, 'num_adj_seats': 3},
                {'name': 'II', 'num_fixed_seats': 0, 'num_adj_seats': 4},
            ],
            'party_vote_info': {
                'name': '-',
                'num_fixed_seats': 0,
                'num_adj_seats': 0,
                'votes': [],
                'specified': False,
                'total': 0,
                'pruned': 0,
            },
            'party_vote_basis': 'totals',
        }
        system = self.make_system(table, 'max-const-seat-share', threshold=50)
        system['adj_determine_divider'] = 'dhondt'
        election = ElectionHandler(
            table, [system], use_thresholds=True).elections[0]
        election.calculate_ref_seat_shares('both')

        self.assertEqual(election.desired_col_sums.tolist(), [7, 0])
        np.testing.assert_allclose(election.fractional_party_seats, [4.2, 2.8])
        np.testing.assert_allclose(
            election.ref_seat_shares.sum(axis=0), [4.2, 2.8])
        np.testing.assert_allclose(election.ref_seat_shares.sum(axis=1), [3, 4])

    def test_fractional_reference_accounts_for_national_seats(self):
        table = load_votes('../data/2-by-2-example.csv')
        table['party_vote_info'] = {
            'name': 'National',
            'num_fixed_seats': 0,
            'num_adj_seats': 2,
            'votes': [4500, 3500],
            'specified': True,
            'total': 8000,
            'pruned': 0,
        }
        table['party_vote_basis'] = 'party_vote_info'
        system = self.make_system(table, 'max-const-seat-share')
        election = ElectionHandler(
            table, [system], use_thresholds=True).elections[0]
        election.calculate_ref_seat_shares('both')

        np.testing.assert_allclose(
            election.total_ref_seat_shares, [15.1875, 11.8125])
        np.testing.assert_allclose(
            election.ref_seat_shares.sum(axis=1), [12, 13])
        self.assertAlmostEqual(election.total_ref_nat.sum(), 2)
        self.assertTrue((election.total_ref_nat >= 0).all())

    def test_simulation_reference_is_independent_of_system_order(self):
        table = load_votes('../data/2-by-2-example.csv')

        def systems():
            first = self.make_system(
                table, 'max-const-seat-share', threshold=50)
            first['name'] = 'Threshold system'
            first['adj_determine_divider'] = 'dhondt'
            second = self.make_system(
                table, 'max-const-seat-share', threshold=0)
            second['name'] = 'No-threshold system'
            second['adj_determine_divider'] = 'sainte-lague'
            return first, second

        settings = SimulationSettings()
        settings['simulation_count'] = 1
        settings['cpu_count'] = 1
        first, second = systems()
        simulations = [
            Simulation(settings, [first, second], table),
            Simulation(settings, list(reversed(systems())), table),
        ]

        references = []
        measures = []
        for simulation in simulations:
            references.append({
                election.system['name']: election.ref_seat_shares.copy()
                for election in simulation.reference_handler.elections
            })
            simulation.run_and_collect_measures(table['votes'], None)
            values = simulation.stat['sum_abs'].mean()
            measures.append(dict(zip(
                [system['name'] for system in simulation.systems], values)))

        for name in references[0]:
            np.testing.assert_allclose(references[0][name], references[1][name])
        np.testing.assert_allclose(
            references[0]['Threshold system'],
            references[0]['No-threshold system'],
        )
        self.assertEqual(measures[0], measures[1])

    def test_representative_methods_preserve_margins(self):
        cases = [
            ('../data/iceland-2021.csv', 'icelandic-law', 5),
            ('../data/norway_2025.csv', 'norwegian-law', 4),
            ('../data/iceland-2021.csv', 'max-const-seat-share', 0),
            ('../data/iceland-2021.csv', 'switching', 0),
            ('../data/iceland-2021.csv', 'swedish-style-switching', 0),
            ('../data/iceland-2021.csv', 'alternating-scaling', 0),
        ]
        for filename, method, threshold in cases:
            with self.subTest(method=method):
                table = load_votes(filename)
                system = self.make_system(table, method, threshold)
                with redirect_stdout(StringIO()):
                    handler = ElectionHandler(table, [deepcopy(system)], True)
                election = handler.elections[0]
                allocation = np.asarray(election.results['all_const_seats'])
                row_totals = [
                    constituency['num_fixed_seats'] + constituency['num_adj_seats']
                    for constituency in election.system['constituencies']
                ]
                self.assertEqual(allocation.sum(axis=1).tolist(), row_totals)
                self.assertEqual(
                    allocation.sum(axis=0).tolist(),
                    election.desired_col_sums.tolist(),
                )
                if method == 'swedish-style-switching':
                    self.assertEqual(len(election.demo_tables), 2)
                    self.assertEqual(
                        election.demo_tables[1]['sup_header'],
                        'Swedish-style switching of adjustment seats',
                    )
