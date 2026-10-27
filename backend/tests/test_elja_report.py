from contextlib import redirect_stdout
import importlib.util
from io import StringIO
import json
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from electionSystem import ElectionSystem
from input_files import load_votes
from sim import write_csv
from simulate import SimulationSettings
from simulation_chunks import combine_chunks, run_chunk, write_chunk_result


SCRIPT = Path(__file__).resolve().parents[1] / 'parsim' / 'elja-report.py'
spec = importlib.util.spec_from_file_location('elja_report', SCRIPT)
report = importlib.util.module_from_spec(spec)
spec.loader.exec_module(report)


class EljaReportTest(unittest.TestCase):
    def make_job(self, directory):
        job = Path(directory) / 'jobs' / 'example-job'
        job.mkdir(parents=True)
        votes = load_votes('../data/2-by-2-example.csv')
        system = ElectionSystem()
        system.copy_info_from_votes(votes)
        system['name'] = 'Test system'
        settings = dict(SimulationSettings(), random_seed=123, simulation_count=5,
                        cpu_count=2, sensitivity=True, sensitivity_covs=[.1],
                        show_additional=True, show_scalings=True, show_single_seat=True)
        display = {'fractional_digits': 5, 'percentage_digits': 2}
        (job / 'inputs.json').write_text(json.dumps({
            'vote_table': votes, 'systems': [system],
            'sim_settings': settings, 'display_settings': display,
        }))
        files = {'Download all file': '/home/user/election.json'}
        (job / 'manifest.json').write_text(json.dumps({
            'job_id': job.name, 'seed': 123, 'replicates': 5, 'first_replicate': 37,
            'nodes': [{'id': 0, 'start': 37, 'replicates': 2},
                      {'id': 3, 'start': 39, 'replicates': 3}],
            'input_files': files,
        }))
        for node, count, start in ((0, 2, 37), (3, 3, 39)):
            _, statistics = combine_chunks(
                [run_chunk(votes, [system], settings, count, start)],
                return_statistics=True)
            write_chunk_result(job / f'node-{node:03d}.json', statistics)
        # Ignore outputs from allocation attempts absent from the manifest.
        (job / 'node-099.json').write_text('unused')
        return job, votes, [system], settings, display, files

    def test_terminology_option_does_not_change_any_job_files(self):
        with TemporaryDirectory() as directory:
            job, *_ = self.make_job(directory)
            before = {path.name: path.read_bytes() for path in job.iterdir()}
            csv_path = Path(directory) / 'district.csv'
            excel_path = Path(directory) / 'district.xlsx'
            process = self.command(job, '-o', csv_path, '-x', excel_path,
                                   '--constituency-term', 'district')
            self.assertEqual(process.returncode, 0, process.stderr)
            self.assertIn('District disparity', csv_path.read_text())
            book = load_workbook(excel_path)
            values = [cell.value for sheet in book for row in sheet for cell in row]
            self.assertIn('District disparity', values)
            book.close()
            self.assertEqual(before, {path.name: path.read_bytes() for path in job.iterdir()})

    def command(self, job, *outputs):
        return subprocess.run([sys.executable, str(SCRIPT), str(job), *map(str, outputs)],
                              capture_output=True, text=True)

    def test_cli_csv_and_excel_match_the_uninterrupted_run_and_preserve_job_files(self):
        with TemporaryDirectory() as directory:
            job, votes, systems, settings, display, files = self.make_job(directory)
            original = {p.name: p.read_bytes() for p in job.iterdir()}
            csv_path, excel_path = Path(directory) / 'out.csv', Path(directory) / 'out.xlsx'
            process = self.command(job, '-o', csv_path, '-x', excel_path)
            self.assertEqual(process.returncode, 0, process.stderr)
            direct = combine_chunks([run_chunk(votes, systems, settings, 5, 37)])
            direct.input_files = files
            expected = Path(directory) / 'expected.csv'
            write_csv(expected, direct, display)
            self.assertEqual(csv_path.read_bytes(), expected.read_bytes())
            self.assertNotIn(b'\r', csv_path.read_bytes())
            self.assertEqual({p.name: p.read_bytes() for p in job.iterdir()}, original)
            workbook = load_workbook(excel_path)
            self.assertIn('Test system', workbook.sheetnames)
            quality = workbook['Quality measures']
            row = next(row for row in quality
                       if row[1].value == 'Loosemore-Hanby index for lists (%)')
            self.assertAlmostEqual(row[2].value, direct.data[0]['lh_lists']['avg'])
            self.assertIn('.00%', row[2].number_format)
            system_sheet = workbook['Test system']
            header = next(cell for row in system_sheet for cell in row
                          if cell.value == 'Double scaled reference seat shares')
            self.assertIn('.00000', system_sheet.cell(header.row + 2, header.column).number_format)
            workbook.close()
            self.assertIn('Replicates: 5; seed: 123', process.stdout)

    def test_job_id_and_excel_only_output(self):
        with TemporaryDirectory() as directory:
            job, *_ = self.make_job(directory)
            output = Path(directory) / 'out.xlsx'
            with patch.object(report, 'HERE', Path(directory)), redirect_stdout(StringIO()):
                self.assertEqual(report.main([job.name, '--excel', str(output)]), 0)
            self.assertTrue(output.exists())

    def test_missing_node_or_wrong_seed_does_not_write_a_report(self):
        for defect in ('missing', 'seed', 'range', 'total', 'duplicate', 'old-format'):
            with self.subTest(defect=defect), TemporaryDirectory() as directory:
                job, *_ = self.make_job(directory)
                node_path = job / 'node-003.json'
                record = json.loads(node_path.read_text())
                manifest_path = job / 'manifest.json'
                manifest = json.loads(manifest_path.read_text())
                if defect == 'missing':
                    node_path.unlink()
                elif defect == 'seed':
                    record['result']['random_seed'] = 999
                elif defect == 'range':
                    record['result']['next_global_iteration'] += 1
                elif defect == 'total':
                    manifest['replicates'] += 1
                elif defect == 'duplicate':
                    manifest['nodes'][1]['id'] = 0
                else:
                    record['format_version'] = 1
                if defect != 'missing':
                    node_path.write_text(json.dumps(record))
                manifest_path.write_text(json.dumps(manifest))
                output = Path(directory) / 'out.csv'
                process = self.command(job, '-o', output)
                self.assertEqual(process.returncode, 1)
                self.assertFalse(output.exists())

    def test_report_paths_cannot_overwrite_the_job_or_each_other(self):
        with TemporaryDirectory() as directory:
            job, *_ = self.make_job(directory)
            for source in ('manifest.json', 'inputs.json', 'node-003.json'):
                path = job / source
                original = path.read_bytes()
                process = self.command(job, '-o', path)
                self.assertEqual(process.returncode, 1)
                self.assertEqual(path.read_bytes(), original)
            path = Path(directory) / 'out'
            process = self.command(job, '-o', path, '-x', path)
            self.assertEqual(process.returncode, 2)
            self.assertFalse(path.exists())


if __name__ == '__main__':
    unittest.main()
