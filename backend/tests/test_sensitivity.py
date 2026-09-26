from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

import entropy_score
from electionSystem import ElectionSystem
from input_util import check_simul_settings
from noweb import load_votes
from sensitivity import (
    generate_perturbations, seat_displacements, sensitivity_covs)
from simulation_excel import SimulationWorkbook
from simulate import Sim_result, Simulation, SimulationSettings


class SensitivityTest(unittest.TestCase):
    def make_system(self, table, name="System-1"):
        system = ElectionSystem()
        system.copy_info_from_votes(table)
        system["name"] = name
        system["adjustment_method"] = "max-const-seat-share"
        return system

    def settings(self, **updates):
        settings = SimulationSettings()
        settings.update({
            "simulation_count": 2,
            "cpu_count": 1,
            "random_seed": 123,
            "sensitivity": True,
            "sensitivity_simulation_count": 3,
            "sensitivity_covs": [1, 2],
        })
        settings.update(updates)
        return settings

    def run_simulation(self, settings=None, systems=None):
        table = load_votes("../data/2-by-2-example.csv")
        systems = systems or [self.make_system(table)]
        simulation = Simulation(settings or self.settings(), systems, table)
        simulation.simulate()
        result = Sim_result(simulation.attributes())
        result.analysis()
        return result

    def test_cov_list_accepts_arbitrary_increasing_values(self):
        self.assertEqual(
            sensitivity_covs([0.1, 0.2, 0.5, 1, 2]),
            [0.001, 0.002, 0.005, 0.01, 0.02],
        )
        with self.assertRaisesRegex(ValueError, "at least one"):
            sensitivity_covs([])
        with self.assertRaisesRegex(ValueError, "positive"):
            sensitivity_covs([0, 1])
        with self.assertRaisesRegex(ValueError, "in increasing order"):
            sensitivity_covs([1, 0.5])

    def test_sensitivity_settings_are_validated(self):
        self.assertEqual(
            SimulationSettings()["sensitivity_simulation_count"], 3)
        self.assertEqual(SimulationSettings()["gen_method"], "log-normal")
        self.assertEqual(SimulationSettings()["sensitivity_gen_method"], "uniform")
        older_settings = SimulationSettings()
        del older_settings["sensitivity_gen_method"]
        self.assertEqual(
            check_simul_settings(older_settings)["sensitivity_gen_method"],
            "uniform")
        settings = self.settings()
        self.assertTrue(check_simul_settings(settings)["sensitivity"])
        settings["sensitivity_simulation_count"] = 0
        with self.assertRaisesRegex(ValueError, "positive integer"):
            check_simul_settings(settings)
        settings["sensitivity"] = False
        self.assertFalse(check_simul_settings(settings)["sensitivity"])

    def test_perturbations_preserve_totals_and_zero_cells(self):
        from randomness import make_rng

        votes = [[100, 0, 50], [0, 40, 60]]
        party_votes = [100, 80, 20]
        perturbed, national = generate_perturbations(
            votes, party_votes, 5, 0.1, "log-normal", make_rng(7, (1,)))

        np.testing.assert_allclose(
            perturbed.sum(axis=2),
            np.tile(np.asarray(votes).sum(axis=1), (5, 1)))
        np.testing.assert_allclose(
            national.sum(axis=1), sum(party_votes))
        self.assertTrue(np.all(perturbed[:, 0, 1] == 0))
        self.assertTrue(np.all(perturbed[:, 1, 0] == 0))

    def test_displacement_is_split_between_and_within_parties(self):
        def election(constituencies, totals):
            return SimpleNamespace(
                party_vote_info={"specified": False},
                results={
                    "all_const_seats": constituencies,
                    "all_grand_total": totals,
                },
            )

        base = election([[2, 0], [0, 2]], [2, 2])
        between = election([[1, 1], [0, 2]], [1, 3])
        within = election([[1, 0], [1, 2]], [2, 2])

        np.testing.assert_array_equal(
            seat_displacements([base], [between]), ([1], [0]))
        np.testing.assert_array_equal(
            seat_displacements([base], [within]), ([0], [1]))

    def test_sensitivity_keeps_ordinary_measures_and_averages_by_major(self):
        result = self.run_simulation()

        self.assertEqual(result.stat["sum_abs"].n, 2)
        self.assertEqual(result.stat["sensitivity_between_parties"].n, 2)
        self.assertEqual(
            result.stat["sensitivity_between_parties"].numpy_mean().shape,
            (2, 1),
        )
        self.assertEqual(result.sensitivity_data["covs"], [1, 2])

    def test_sensitivity_combines_reproducibly_across_workers(self):
        table = load_votes("../data/2-by-2-example.csv")
        systems = [self.make_system(table)]

        def run(count, start):
            settings = self.settings(simulation_count=count)
            simulation = Simulation(
                settings, deepcopy(systems), deepcopy(table),
                start_iteration=start)
            simulation.simulate()
            return Sim_result(simulation.attributes())

        uninterrupted = run(4, 0)
        combined = run(2, 0)
        combined.combine(run(2, 2))
        for measure in (
                "sensitivity_between_parties",
                "sensitivity_within_parties"):
            np.testing.assert_allclose(
                uninterrupted.stat[measure].numpy_mean(),
                combined.stat[measure].numpy_mean())
            np.testing.assert_allclose(
                uninterrupted.stat[measure].numpy_std(),
                combined.stat[measure].numpy_std())

    def test_entropy_reference_is_not_computed_for_minor_perturbations(self):
        settings = self.settings(entropy_score=True, simulation_count=2)
        with patch.object(
                entropy_score,
                "_optimal_allocation",
                wraps=entropy_score._optimal_allocation) as optimize:
            self.run_simulation(settings)
        self.assertEqual(optimize.call_count, 2)

    def test_sensitivity_is_last_quality_block_on_web_and_in_excel(self):
        result = self.run_simulation()
        web = result.get_result_web(False)
        self.assertEqual(web["vuedata"]["group_ids"][-2:], [
            "sensitivityWithin", "sensitivityBetween"])
        for group in (
                "cmpListTitle", "cmpList", "cmpPartyTitle", "cmpParty",
                "sensitivityWithin", "sensitivityBetween"):
            self.assertEqual(web["vuedata"]["group_stats"][group], ["avg"])
        for group in ("shareTitle", "toLists", "toPartiesTotal", "other"):
            self.assertNotIn(group, web["vuedata"]["group_stats"])
        self.assertEqual(
            [row["rowtitle"]
             for row in web["vuedata"]["sensitivityWithin"]],
            ["1% CoV", "2% CoV"],
        )

        with TemporaryDirectory() as directory:
            path = Path(directory) / "sensitivity.xlsx"
            SimulationWorkbook(web, path).write()
            sheet = load_workbook(path, read_only=True)["Quality measures"]
            first_column = [
                row[0].value for row in sheet.iter_rows(min_col=1, max_col=1)]
        self.assertLess(
            first_column.index("Seats displaced between lists within parties"),
            first_column.index("Seats displaced between parties"),
        )
        self.assertEqual(first_column.count("1% CoV"), 2)
        self.assertEqual(first_column.count("2% CoV"), 2)


if __name__ == "__main__":
    unittest.main()
