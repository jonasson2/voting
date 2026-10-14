"""Regression cases for geographical results and regional entropy benchmarks."""

import csv
from itertools import combinations_with_replacement
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import numpy as np
from openpyxl import load_workbook

from electionHandler import ElectionHandler
from electionSystem import ElectionSystem
from entropy_score import _problem_key, calculate
from excel_util import simulation_to_xlsx
from methods.optimal_lp import optimal_lp
from division_rules import dhondt_gen
from noweb import load_votes
from reference_measures import adjustment_spread, geographical_indices, local_indices
from sim import write_csv
from simulate import Collect, Simulation, SimulationSettings
from simulation_chunks import (
    combine_chunks, read_chunk_result, run_chunk, write_chunk_result)
from table_util import entropy


def regional_example():
    votes = load_votes('../data/2-by-2-example.csv')
    votes['votes'] = [[100, 60], [90, 50]]
    for index, constituency in enumerate(votes['constituencies']):
        constituency.update(num_fixed_seats=0, num_adj_seats=1,
                            region=str(index))
    votes['regions'] = [
        {'abbreviation': str(index), 'name': str(index), 'num_adj_seats': 1}
        for index in range(2)
    ]
    system = ElectionSystem()
    system.copy_info_from_votes(votes)
    system['adjustment_method'] = 'max-const-votes'
    return votes, system


class GeographicalMeasuresTest(unittest.TestCase):
    def test_local_indices_weight_units_by_seats_and_aggregate_regions_first(self):
        votes = [[50, 30, 20], [100, 0, 0]]
        seats = [[12, 5, 3], [0, 2, 0]]
        lh, gallagher = local_indices(votes, seats)
        self.assertAlmostEqual(lh, (20 * .1 + 2) / 22)
        self.assertAlmostEqual(gallagher, (20 * np.sqrt(.0075) + 2) / 22)
        # National geographical shares match, but local party shares do not.
        votes = [[60, 40], [40, 60]]
        seats = [[1, 0], [0, 1]]
        self.assertEqual(geographical_indices([100, 100], [1, 1]), (0, 0))
        np.testing.assert_allclose(local_indices(votes, seats), [.4, .4])
        self.assertEqual(local_indices(votes, seats, [[0, 1]]), (0, 0))
        self.assertEqual(local_indices([[0, 0]], [[1, 0]]), (0, 0))
        self.assertEqual(local_indices([[1, 0]], [[0, 0]]), (0, 0))
        # A zero-seat unit has zero weight; a positive-vote zero-seat party
        # in an occupied unit remains part of the discrepancy.
        np.testing.assert_allclose(
            local_indices([[60, 20, 20], [0, 100, 0]], [[1, 0, 0], [0, 0, 0]]),
            [.4, np.sqrt(.12)])

    def test_indices_normalization_cancellation_and_zero_totals(self):
        lh, gallagher = geographical_indices([60, 40], [55, 45])
        self.assertAlmostEqual(lh, .05)
        self.assertAlmostEqual(gallagher, .05)
        lh, gallagher = geographical_indices([50, 30, 20], [60, 25, 15])
        self.assertAlmostEqual(lh, .1)
        self.assertAlmostEqual(gallagher, .08660254037844386)
        self.assertEqual(geographical_indices([0, 0], [1, 0]), (0, 0))
        self.assertEqual(geographical_indices([1, 0], [0, 0]), (0, 0))
        self.assertEqual(geographical_indices([1, 0], [1, 0]), (0, 0))
        self.assertEqual(geographical_indices([1, 0], [0, 1]), (1, 1))
        votes = np.array([25, 25, 25, 25])
        seats = np.array([30, 20, 30, 20])
        self.assertGreater(geographical_indices(votes, seats)[0], 0)
        groups = [[0, 1], [2, 3]]
        self.assertEqual(geographical_indices(
            [votes[g].sum() for g in groups],
            [seats[g].sum() for g in groups]), (0, 0))

    def test_spread_counts_recipients_and_zero_weight(self):
        adjustment = [[1, 1, 0], [2, 0, 0], [0, 0, 0]]
        seats = [6, 3, 1]
        self.assertEqual(adjustment_spread(adjustment, seats), 1.5)
        self.assertEqual(adjustment_spread(adjustment, seats, [[0, 1], [2]]), .9)
        adjustment[1], adjustment[2] = adjustment[2], adjustment[1]
        self.assertEqual(adjustment_spread(adjustment, seats), 1.5)
        self.assertEqual(adjustment_spread(adjustment, seats, [[0, 1], [2]]), 1.5)
        self.assertEqual(adjustment_spread([[0, 0]], [6, 3]), 0)
        self.assertEqual(adjustment_spread([[1, 0]], [0, 0]), 0)

    def test_regional_entropy_uses_global_optimum_and_cache_constraints(self):
        votes, system = regional_example()
        election = ElectionHandler(votes, [system], True).elections[0]
        self.assertEqual(election.results['all_const_seats'], [[1, 0], [0, 1]])
        cache = {}
        self.assertAlmostEqual(calculate(election, cache), 5000 / 5400)
        key = _problem_key(election)
        election.region_party_totals = np.array([[0, 1], [1, 0]])
        self.assertEqual(_problem_key(election), key)
        self.assertAlmostEqual(calculate(election, cache), 5000 / 5400)
        election.regions[0]['num_adj_seats'] = 0
        election.regions[1]['num_adj_seats'] = 2
        self.assertNotEqual(_problem_key(election), key)

        votes, system = regional_example()
        system['adjustment_method'] = 'optimal-lp'
        election = ElectionHandler(votes, [system], True).elections[0]
        # The within-region optimum is still globally suboptimal when the
        # intermediate party-by-region assignment is greedy.
        self.assertAlmostEqual(calculate(election), 5000 / 5400)
        system['regional_adjustment_method'] = 'optimal-lp'
        election = ElectionHandler(votes, [system], True).elections[0]
        self.assertEqual(election.results['all_const_seats'], [[0, 1], [1, 0]])
        self.assertEqual(calculate(election), 1)

    def test_regional_budgets_hold_with_flexible_constituency_bounds(self):
        votes = [[100, 10], [90, 10], [80, 10], [70, 10]]
        prior = np.zeros((4, 2), dtype=int)
        allocation, _ = optimal_lp(
            votes, [0] * 4, [1, 1], prior, dhondt_gen,
            num_adjustment_seats=2, min_adj_seats=[0] * 4,
            max_adj_seats=[None] * 4,
            region_groups=[[0, 1], [2, 3]], region_adj_seats=[1, 1],
        )
        self.assertEqual(allocation[[0, 1]].sum(), 1)
        self.assertEqual(allocation[[2, 3]].sum(), 1)
        np.testing.assert_array_equal(allocation.sum(axis=0), [1, 1])
        feasible = []
        for chosen in combinations_with_replacement(range(8), 2):
            candidate = np.bincount(chosen, minlength=8).reshape(4, 2)
            if (candidate[[0, 1]].sum() == candidate[[2, 3]].sum() == 1
                    and np.array_equal(candidate.sum(axis=0), [1, 1])):
                feasible.append(candidate)
        self.assertAlmostEqual(
            entropy(votes, allocation, dhondt_gen),
            max(entropy(votes, candidate, dhondt_gen) for candidate in feasible))

        prior[0, 0] = 1
        allocation, _ = optimal_lp(
            votes, prior.sum(axis=1), [3, 1], prior, dhondt_gen,
            nat_prior_allocations=[0, 0],
            num_adjustment_seats=2, min_adj_seats=[0] * 4,
            max_adj_seats=[None] * 4,
            region_groups=[[0, 1], [2, 3]], region_adj_seats=[1, 1],
        )
        self.assertGreaterEqual(allocation[0, 0], 1)
        self.assertEqual(allocation[[0, 1]].sum(), 2)
        self.assertEqual(allocation[[2, 3]].sum(), 1)
        self.assertTrue(np.all(allocation.sum(axis=0) <= [3, 1]))

    def test_geographical_votes_include_pruned_parties(self):
        votes, system = regional_example()
        votes['pruned'] = [20, 0]
        simulation = Simulation(SimulationSettings(), [system], votes)
        election = simulation.election_handler.elections[0]
        election.calculate_ref_seat_shares(simulation.primary_scaling)
        deviations = Collect()
        simulation.reference_measures(election, deviations)
        expected = geographical_indices([180, 140], [1, 1])
        self.assertAlmostEqual(deviations['lh_constituencies'][0], expected[0])
        self.assertAlmostEqual(deviations['gallagher_lists'][0],
                               (60 / 160 + 90 / 140) / 2)

    def test_one_region_measures_party_proportionality_not_geography(self):
        votes, system = regional_example()
        votes['regions'] = [
            {'abbreviation': 'all', 'name': 'All', 'num_adj_seats': 2}]
        for constituency in votes['constituencies']:
            constituency['region'] = 'all'
        settings = SimulationSettings()
        settings.update(random_seed=41, simulation_count=1, cpu_count=1,
                        const_rsd=0)
        result = combine_chunks([run_chunk(votes, [system], settings, 1, 0)])
        self.assertAlmostEqual(result.data[0]['lh_region_parties']['avg'], 2 / 15)
        self.assertAlmostEqual(result.data[0]['gallagher_region_parties']['avg'], 2 / 15)

    def test_all_reports_and_chunk_merge(self):
        settings = SimulationSettings()
        settings.update(random_seed=41, simulation_count=4, cpu_count=1)
        votes, system = regional_example()
        direct = combine_chunks([run_chunk(votes, [system], settings, 4, 0)])
        merged = combine_chunks([
            run_chunk(votes, [system], settings, 2, 0),
            run_chunk(votes, [system], settings, 2, 2),
        ])
        new_measures = ('lh_region_parties', 'gallagher_region_parties',
                        'gallagher_lists',
                        'adjustment_regional_spread',
                        'adjustment_constituency_spread')
        for measure in new_measures:
            for statistic in ('avg', 'std'):
                self.assertAlmostEqual(direct.data[0][measure][statistic],
                                       merged.data[0][measure][statistic])
        report = merged.get_result_web(False)
        local_measures = {row['measure'] for row in report['vuedata']['const']}
        self.assertTrue({'gallagher_lists', 'lh_region_parties',
                         'gallagher_region_parties'} <= local_measures)
        self.assertNotIn('gallagher_constituencies', direct.MEASURES)
        self.assertNotIn('lh_regions', direct.MEASURES)
        self.assertNotIn('gallagher_regions', direct.MEASURES)
        measures = {row['measure'] for group in ('const', 'parity')
                    for row in report['vuedata'][group]}
        self.assertTrue(set(new_measures) <= measures)
        with TemporaryDirectory() as directory:
            csv_path = Path(directory) / 'result.csv'
            xlsx_path = Path(directory) / 'result.xlsx'
            chunk_path = Path(directory) / 'chunk.json'
            write_chunk_result(chunk_path, run_chunk(votes, [system], settings, 2, 0))
            restored = combine_chunks([
                read_chunk_result(chunk_path),
                run_chunk(votes, [system], settings, 2, 2),
            ])
            self.assertEqual(restored.regional_measure_available, [True])
            for measure in new_measures:
                self.assertAlmostEqual(restored.data[0][measure]['avg'],
                                       direct.data[0][measure]['avg'])
            write_csv(csv_path, merged)
            with csv_path.open() as stream:
                csv_rows = list(csv.reader(stream))
                labels = [row[1] for row in csv_rows if len(row) > 1]
            self.assertIn('Gallagher index for regions (%)', labels)
            percentage = next(row[2] for row in csv_rows
                              if len(row) > 2 and row[1] == 'Gallagher index for regions (%)')
            self.assertTrue(percentage.endswith(')%'), percentage)
            simulation_to_xlsx(report, str(xlsx_path))
            workbook = load_workbook(xlsx_path)
            excel_rows = [row for row in workbook['Quality measures']]
            labels = [row[1].value for row in excel_rows]
            workbook.close()
            self.assertIn('Adjustment-seat spread across regions', labels)
            region_index = labels.index('Gallagher index for regions (%)')
            self.assertIn('%', excel_rows[region_index][2].number_format)

        votes.pop('regions')
        for constituency in votes['constituencies']:
            constituency.pop('region')
        plain = combine_chunks([run_chunk(votes, [system], settings, 1, 0)])
        report = plain.get_result_web(False)
        measures = {row['measure'] for group in ('const', 'parity')
                    for row in report['vuedata'][group]}
        self.assertIn('gallagher_lists', measures)
        self.assertNotIn('gallagher_region_parties', measures)
        self.assertNotIn('adjustment_regional_spread', measures)
        with TemporaryDirectory() as directory:
            csv_path = Path(directory) / 'plain.csv'
            xlsx_path = Path(directory) / 'plain.xlsx'
            write_csv(csv_path, plain)
            self.assertNotIn('Gallagher index for regions', csv_path.read_text())
            simulation_to_xlsx(plain.get_result_web(False), str(xlsx_path))
            workbook = load_workbook(xlsx_path)
            self.assertFalse(any(
                cell.value == 'Adjustment-seat spread across regions'
                for row in workbook['Quality measures'] for cell in row))
            workbook.close()


if __name__ == '__main__':
    unittest.main()
