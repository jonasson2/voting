from copy import deepcopy
import csv
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import numpy as np

from electionSystem import ElectionSystem
from electionHandler import ElectionHandler
from input_files import load_votes, validate_settings
from reference_measures import SCALINGS, list_measures, selected_scalings
from running_stats import Running_stats
from simulate import SimulationSettings
from simulation_chunks import run_chunk, combine_chunks, write_chunk_result, read_chunk_result
from sim import write_csv
from voting import Election


class ReferenceMeasuresTest(unittest.TestCase):
    def test_double_reference_matrices_merge_and_export_beside_local_matrices(self):
        from openpyxl import load_workbook
        from excel_util import simulation_to_xlsx
        for national_votes in (False, True):
            with self.subTest(national_votes=national_votes):
                votes, systems, settings = self.inputs()
                settings['sensitivity'] = False
                if national_votes:
                    votes['party_vote_info'].update(
                        specified=True, name='National', votes=[4500, 3500],
                        total=8000, num_fixed_seats=0, num_adj_seats=2)
                    votes['party_vote_basis'] = 'party_vote_info'
                    systems[0]['adjustment_method'] = 'max-const-seat-share'
                chunks = [run_chunk(votes, systems, settings, 1, index)
                          for index in range(4)]
                matrices = np.array([
                    Running_stats.from_dict(deepcopy(
                        chunk['stat']['double_ref_seat_shares'][0])).mean()
                    for chunk in chunks])
                merged = combine_chunks(deepcopy(chunks))
                direct = combine_chunks([run_chunk(votes, systems, settings, 4, 0)])
                statistics = merged.seat_data[0]['double_ref_seat_shares']
                np.testing.assert_allclose(statistics['avg'], matrices.mean(axis=0))
                np.testing.assert_allclose(statistics['std'], matrices.std(axis=0, ddof=1), atol=1e-12)
                np.testing.assert_allclose(statistics['avg'],
                                           direct.seat_data[0]['double_ref_seat_shares']['avg'])
                np.testing.assert_allclose(np.array(statistics['avg'])[:, -1],
                    np.array(merged.seat_data[0]['total_seats']['avg'])[:, -1])
                if national_votes:
                    # The grand totals use integer national targets; the
                    # national row supplies the part outside constituencies.
                    np.testing.assert_allclose(matrices[:, -1, :-1],
                        matrices[:, -3, :-1] + matrices[:, -2, :-1])
                with TemporaryDirectory() as directory:
                    path = Path(directory) / 'double-reference.xlsx'
                    simulation_to_xlsx(merged.get_result_web(False), str(path))
                    workbook = load_workbook(path)
                    sheet = workbook[systems[0]['name'][:31]]
                    header = next(cell for row in sheet for cell in row
                                  if cell.value == 'Double scaled reference seat shares')
                    local_header = next(cell for row in sheet for cell in row
                                        if cell.value == 'Reference seat shares')
                    self.assertEqual(header.row, local_header.row)
                    self.assertEqual(header.column - local_header.column,
                                     len(votes['parties']) + 2)
                    for label, matrix in (
                            ('Values based on source votes', merged.base_allocations[0]['double_ref_seat_shares']),
                            ('Avg. simulated values', statistics['avg']),
                            ('Standard deviations', statistics['std'])):
                        top = next(cell.row for row in sheet for cell in row[:1]
                                   if cell.value == label)
                        exported = [[cell.value for cell in row] for row in sheet.iter_rows(
                            min_row=top, max_row=top + len(matrix) - 1,
                            min_col=header.column,
                            max_col=header.column + len(votes['parties']))]
                        np.testing.assert_allclose(exported, matrix, atol=1e-8)
                    workbook.close()

    def inputs(self):
        votes = load_votes('../data/2-by-2-example.csv')
        system = ElectionSystem()
        system.copy_info_from_votes(votes)
        settings = SimulationSettings()
        settings.update(random_seed=317, simulation_count=4, cpu_count=1,
                        sensitivity=True, sensitivity_simulation_count=2)
        return votes, [system], settings

    def test_settings_migrate_old_scalings_to_display_option(self):
        for choice, expected in [('both', True), (['party', 'const', 'party'], True),
                                 ('const', False), ([], False)]:
            settings = dict(SimulationSettings(), scaling=choice)
            del settings['show_scalings']
            validated = validate_settings(settings)
            self.assertNotIn('scaling', validated)
            self.assertEqual(validated['show_scalings'], expected)
        settings = dict(SimulationSettings(), scaling='both', show_scalings='false')
        self.assertFalse(validate_settings(settings)['show_scalings'])
        self.assertFalse(validate_settings(SimulationSettings())['show_scalings'])
        self.assertTrue(validate_settings(dict(SimulationSettings(), show_scalings='true'))['show_scalings'])
        for choice in [None, 'unknown', ['both', 'bad'], [False], {'const': True}]:
            with self.assertRaises(ValueError):
                selected_scalings(choice)

    def test_known_discrepancies_and_zero_reference_penalties(self):
        values = list_measures([[2, 0], [0, 2]], [[1, 1], [1, 1]])
        self.assertEqual(values['deviation'], 4)
        self.assertEqual(values['surplus'], 1)
        self.assertEqual(values['shortfall'], 1)
        self.assertEqual(values['squared'], 4)
        self.assertEqual(values['relative_squared'], 4)
        self.assertEqual(values['share_surplus'], .5)
        self.assertEqual(values['share_shortfall'], .5)
        self.assertEqual(values['overrepresentation'], 2)
        self.assertEqual(values['underrepresentation'], 1)
        self.assertEqual(values['relative_absolute'], 4)
        zero = list_measures([[0, 1]], [[0, 1]])
        self.assertEqual(zero['relative_squared'], 0)
        penalty = list_measures([[1, 0]], [[0, 1]])
        self.assertEqual(penalty['relative_squared'], float('inf'))
        self.assertEqual(penalty['overrepresentation'], float('inf'))

    def test_integer_party_reference_follows_national_divisor_rounding(self):
        table, _, _ = self.inputs()
        table['parties'] = ['A', 'B', 'C']
        table['votes'] = [[50, 30, 20], [50, 30, 20]]
        for constituency, seats in zip(table['constituencies'], [1, 2]):
            constituency.update(num_fixed_seats=0, num_adj_seats=seats)
        for divider, targets in [('dhondt', [2, 1, 0]),
                                 ('sainte-lague', [1, 1, 1])]:
            system = ElectionSystem()
            system.copy_info_from_votes(table)
            system.update(adjustment_method='max-const-seat-share',
                          adj_determine_divider=divider, adjustment_threshold=0)
            election = ElectionHandler(table, [system], use_thresholds=True).elections[0]
            with self.subTest(divider=divider):
                np.testing.assert_array_equal(election.desired_col_sums, targets)
                for scaling in ('party', 'both'):
                    reference = election.reference_seats(scaling)
                    np.testing.assert_allclose(reference.sum(axis=0), targets)
                    if scaling == 'both':
                        np.testing.assert_allclose(reference.sum(axis=1), [1, 2])
                np.testing.assert_allclose(election.reference_seats('party'),
                                          np.array([targets, targets]) / 2)
                np.testing.assert_allclose(election.reference_seats('total'),
                                          [[.75, .45, .3], [.75, .45, .3]])

    def test_integer_party_reference_preserves_fixed_seat_overhang(self):
        table, _, _ = self.inputs()
        table['votes'] = [[1, 99], [999, 1]]
        table['constituencies'][0].update(num_fixed_seats=3, num_adj_seats=0)
        table['constituencies'][1].update(num_fixed_seats=0, num_adj_seats=1)
        system = ElectionSystem()
        system.copy_info_from_votes(table)
        system.update(adjustment_method='max-const-seat-share',
                      primary_divider='dhondt', adj_determine_divider='dhondt',
                      adjustment_threshold=0)
        election = ElectionHandler(table, [system], use_thresholds=True).elections[0]
        np.testing.assert_array_equal(election.desired_col_sums, [1, 3])
        for scaling in ('party', 'both'):
            with self.subTest(scaling=scaling):
                np.testing.assert_allclose(
                    election.reference_seats(scaling).sum(axis=0), [1, 3])
        np.testing.assert_allclose(election.reference_seats('total').sum(axis=0),
                                  [40 / 11, 4 / 11])

    def test_all_benchmarks_computed_independently_of_display_and_legacy_choices(self):
        votes, systems, settings = self.inputs()
        cases = []
        for options in [dict(show_scalings=False), dict(show_scalings=True),
                        dict(scaling=[]), dict(scaling=['both'])]:
            case_settings = dict(settings, **options)
            calls = []
            original = Election.reference_seats
            def record(election, scaling):
                calls.append(scaling)
                return original(election, scaling)
            with patch.object(Election, 'reference_seats', record):
                result = combine_chunks([run_chunk(votes, systems, case_settings, 4, 0)])
            self.assertEqual(set(calls), set(SCALINGS))
            self.assertEqual(set(key.split('_')[0] for key in result.MEASURES
                                 if key.endswith('_deviation')), set(SCALINGS))
            self.assertNotIn('scaling', result.sim_settings)
            self.assertEqual(result.primary_scaling, 'const')
            cases.append(result)
        for result in cases[1:]:
            for measure in result.MEASURES:
                np.testing.assert_equal(result.stat[measure].mean(),
                                        cases[0].stat[measure].mean())
            np.testing.assert_equal(result.sensitivity_data, cases[0].sensitivity_data)
            np.testing.assert_equal(result.stat['total_seats'][0].mean(),
                                    cases[0].stat['total_seats'][0].mean())
        reference = cases[-1]
        seats = sum(c['num_fixed_seats'] + c['num_adj_seats'] for c in votes['constituencies'])
        self.assertAlmostEqual(reference.data[0]['lh_lists']['avg'],
                               reference.data[0]['const_deviation']['avg'] / (2 * seats))

    def test_multiple_benchmarks_merge_and_csv_options(self):
        votes, systems, settings = self.inputs()
        direct = combine_chunks([run_chunk(votes, systems, settings, 4, 0)])
        with TemporaryDirectory() as directory:
            paths = [Path(directory) / f'{index}.json' for index in range(2)]
            for index, path in enumerate(paths):
                write_chunk_result(path, run_chunk(votes, systems, settings, 2, index * 2))
            merged = combine_chunks([read_chunk_result(path) for path in paths])
            for measure in direct.MEASURES:
                for stat in ('avg', 'std'):
                    self.assertAlmostEqual(direct.data[0][measure][stat],
                                           merged.data[0][measure][stat], places=10)
            output = Path(directory) / 'results.csv'
            for show in (False, True):
                for show_scalings in (False, True):
                    merged.sim_settings.update(show_additional=show, show_scalings=show_scalings,
                                               show_single_seat=show)
                    write_csv(output, merged)
                    with output.open() as file:
                        rows = list(csv.reader(file))
                    labels = [row[1] for row in rows if len(row) > 1]
                    headings = [row[0] for row in rows]
                    self.assertEqual(headings.count('Additional proportionality measures'), int(show))
                    self.assertEqual(headings.count('Double, party, and national scaling'), int(show_scalings))
                    self.assertEqual(headings.count('Specifically for single-seat constituencies'), int(show))
                    self.assertEqual('Total squared list seat deviation (Hare quota)' in labels, show)
                    self.assertEqual('Potential overhang' in labels, show)
                    self.assertEqual('Total squared party-seat deviation' in labels, show)
                    self.assertEqual(labels.count('Total list seat deviation'), 4 if show_scalings else 1)
                    self.assertEqual(labels.count('Local squared deviation per reference seat'), 1)
                    self.assertEqual(labels.count('Total squared list seat deviation (Hare quota)'),
                                     (4 if show_scalings else 1) if show else 0)

    def test_excel_retains_optional_and_secondary_rows_with_percent_formats(self):
        import openpyxl
        from excel_util import simulation_to_xlsx
        votes, systems, settings = self.inputs()
        settings.update(show_scalings=False, show_additional=False,
                        show_single_seat=False)
        result = combine_chunks([run_chunk(votes, systems, settings, 2, 0)])
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'results.xlsx'
            simulation_to_xlsx(result.get_result_web(False), str(path))
            workbook = openpyxl.load_workbook(path)
            sheet = workbook['Quality measures']
            rows = {row[1].value: row for row in sheet if row[1].value}
            self.assertIn('Total squared list seat deviation (Hare quota)', rows)
            self.assertIn('Potential overhang', rows)
            local_index = rows['Loosemore-Hanby index for lists (%)']
            self.assertIn('%', local_index[2].number_format)
            self.assertIsNotNone(local_index[1].comment)
            labels = [cell.value for row in sheet for cell in row[:2]]
            blocks = [row[0].value for row in sheet]
            for title in ('List allocation quality measures',
                          'Total allocation quality measures',
                          'Additional proportionality measures',
                          'Double, party, and national scaling',
                          'Specifically for single-seat constituencies'):
                self.assertEqual(blocks.count(title), 1)
            local_squared_row = rows['Total squared list seat deviation (Hare quota)'][1].row
            additional_start = blocks.index('Additional proportionality measures') + 1
            scaling_start = blocks.index('Double, party, and national scaling') + 1
            # All four named-method variants belong to the additional block.
            self.assertGreater(local_squared_row, additional_start)
            self.assertLess(local_squared_row, scaling_start)
            self.assertLess(additional_start, scaling_start)
            self.assertIn('Over-allocation per reference seat', labels)
            self.assertIn('Measures with\ndouble scaling', labels)
            self.assertIn('Measures with\nparty scaling', labels)
            self.assertIn('Measures with\nnational scaling', labels)
            workbook.close()

    def test_infinite_penalty_survives_disk_and_merge_without_false_ci(self):
        votes, systems, settings = self.inputs()
        first = run_chunk(votes, systems, settings, 2, 0)
        counter = Running_stats.from_dict(first['stat']['local_squared'])
        counter.update([float('inf')])
        first['stat']['local_squared'] = counter.to_dict()
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'stats.json'
            write_chunk_result(path, first)
            result = combine_chunks([read_chunk_result(path),
                                     run_chunk(votes, systems, settings, 2, 2)])
            self.assertEqual(result.data[0]['local_squared']['avg'], '∞')
            self.assertEqual(result.data[0]['local_squared']['std'], '–')
            table = result.get_result_web(False)['vuedata']
            row = next(row for row in table['const'] if row['measure'] == 'local_squared')
            self.assertEqual(row['avg'], ['∞'])
            json.dumps(table, allow_nan=False)
