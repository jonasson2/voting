from copy import deepcopy
import csv
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

from dictionaries import ELECTION_LAW_PRESETS
from division_rules import dhondt_gen, hare, sainte_lague_gen
from electionHandler import ElectionHandler
from electionSystem import ElectionSystem
from methods import danish
from noweb import load_votes, votes_to_excel
from simulate import Simulation, SimulationSettings, simulation_vote_table
from vote_table import check_vote_table, process_vote_table
from web import app

DATA = Path(__file__).resolve().parents[2] / "data"


def system_for(table):
    system = ElectionSystem()
    system.copy_info_from_votes(table)
    settings = next(p["settings"] for p in ELECTION_LAW_PRESETS if p["value"] == "denmark")
    system.update({key: value for key, value in settings.items()
                   if key != "constituency_seat_specification"})
    return system


class DanishTest(unittest.TestCase):
    def setUp(self):
        self.table = load_votes(DATA / "denmark_2026.csv")
        self.rng = np.random.default_rng(123)

    def election(self, table=None):
        table = table or self.table
        return ElectionHandler(table, [system_for(table)], True).elections[0]

    def test_official_2026_all_constituency_and_regional_results(self):
        expected = json.loads((DATA / "denmark/official-results_2026.json").read_text())
        election = self.election()
        for c, constituency in enumerate(expected["constituencies"]):
            for label, result in constituency["parties"].items():
                if label == "Uden for partierne":
                    continue
                p = self.table["parties"].index(label.split(". ", 1)[0])
                for field, key in [("fixed_seats", "fixed_const_seats"),
                                   ("total_seats", "all_const_seats"),
                                   ("adjustment_seats", "adj_const_seats")]:
                    self.assertEqual(election.results[key][c][p], result[field], (constituency["name"], label, field))
        self.assertEqual(sum(election.results["all_const_total"]), 175)
        self.assertEqual(election.results["all_const_total"][12:], [0] * 6)
        for r, region in enumerate(expected["regions"]):
            for label, result in region["parties"].items():
                if label != "Uden for partierne":
                    p = self.table["parties"].index(label.split(". ", 1)[0])
                    self.assertEqual(election.region_party_totals[r, p], result["total_seats"])

    def test_step_tables_and_excel(self):
        handler = ElectionHandler(self.table, [system_for(self.table)], True)
        election = handler.elections[0]
        self.assertEqual(len(election.demo_tables), 2)
        for table in election.demo_tables:
            self.assertEqual([row[0] for row in table["steps"]], list(range(1, 41)))
            self.assertEqual(len(table["format"]), len(table["headers"]))
            for row in table["steps"]:
                self.assertAlmostEqual(row[-2], row[-4] / row[-3])
        with TemporaryDirectory() as directory:
            path = Path(directory) / "election.xlsx"
            handler.to_xlsx(path)
            book = load_workbook(path)
            strings = [cell.value for sheet in book for row in sheet for cell in row]
            for table in election.demo_tables:
                self.assertIn(table["sup_header"], strings)
            book.close()

    def test_vote_excel_and_json_round_trips(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "votes.xlsx"
            votes_to_excel(self.table, path)
            self.assertEqual(load_votes(path), self.table)
        self.assertEqual(check_vote_table(json.loads(json.dumps(self.table))), self.table)

    def test_regions_with_exact_adjustment_seats_and_no_names(self):
        rows = list(csv.reader([
            "Example,fixed,adj,region,A,B",
            "Independent candidates,,,,0,0",
            "North,1,1,H,90,10",
            "South,1,1,SS,40,60",
            ",,,,,",
            "Regions,Name,adj,,,",
            "H,Capital,1,,,",
            "SS,South,1,,,",
        ]))
        table = check_vote_table(process_vote_table(rows, "test.csv"))
        election = self.election(table)
        self.assertEqual(sum(election.results["all_const_total"]), 4)
        self.assertEqual(election.final_row_sums.tolist(), [2, 2])
        with TemporaryDirectory() as directory:
            path = Path(directory) / "votes.xlsx"
            votes_to_excel(table, path)
            reloaded = load_votes(path)
            self.assertEqual(reloaded["regions"], table["regions"])
            self.assertEqual(reloaded["constituencies"], table["constituencies"])

    def test_constituency_bounds_are_enforced_inside_regions(self):
        table = deepcopy(self.table)
        table["constituencies"][0]["num_adj_seats"] = 4
        table["constituencies"][0]["max_adj_seats"] = 4
        election = self.election(check_vote_table(table))
        self.assertEqual(sum(election.results["adj_const_seats"][0]), 4)
        self.assertEqual(sum(election.results["adj_const_total"]), 40)

    def test_upload_regions_and_independents(self):
        with app.test_client() as client:
            response = client.post("/api/votes/upload/", data={
                "file": (BytesIO((DATA / "denmark_2026.csv").read_bytes()), "votes.csv")})
        self.assertNotIn("error", response.json)
        self.assertEqual(response.json["regions"], self.table["regions"])
        self.assertEqual(response.json["independent_candidates"], [False] * 12 + [True] * 6)

    def test_malformed_regional_metadata(self):
        with (DATA / "denmark_2026.csv").open() as file:
            rows = list(csv.reader(file))
        for mutation in (
            lambda r: r[2].__setitem__(5, "2"),
            lambda r: r[-1].__setitem__(0, "H"),
            lambda r: r[-1].__setitem__(2, "15"),
            lambda r: r[4].__setitem__(4, "unknown"),
            lambda r: r[-2].__setitem__(4, "unexpected"),
            lambda r: r.insert(2, r[2].copy()),
        ):
            altered = deepcopy(rows)
            mutation(altered)
            self.assertIsInstance(process_vote_table(altered, "test.csv"), str)

    def test_validation_of_national_votes_regions_and_independent_flags(self):
        table = deepcopy(self.table)
        table["party_vote_info"].update(specified=True, votes=[1] * 18)
        with self.assertRaisesRegex(ValueError, "National party votes"):
            check_vote_table(table)
        for mutate in (
            lambda t: t["independent_candidates"].pop(),
            lambda t: t["independent_candidates"].__setitem__(0, 1),
            lambda t: t["constituencies"][0].__setitem__("region", "ZZ"),
            lambda t: t["regions"][0].__setitem__("num_adj_seats", 13),
        ):
            table = deepcopy(self.table)
            mutate(table)
            with self.assertRaises(ValueError):
                check_vote_table(table)

    def test_all_eligibility_routes_and_denominators(self):
        votes = np.array([[1, 10, 200, 1, 20], [1, 10, 200, 1, 20], [1, 0, 200, 1, 20]])
        fixed = np.array([[1, 0, 9, 0, 0], [0, 0, 10, 0, 0], [0, 0, 10, 0, 0]])
        groups = [np.array([c]) for c in range(3)]
        eligible = danish.eligible_parties(votes, fixed, np.array([0, 0, 0, 0, 1], bool),
            groups, np.array([100, 100, 2000]), threshold=2)
        np.testing.assert_array_equal(eligible, [True, True, True, False, False])
        # Adding retained/pruned votes defeats the regional and national tests.
        eligible = danish.eligible_parties(votes, fixed, np.array([0, 0, 0, 0, 1], bool),
            groups, np.array([1000, 1000, 2000]), threshold=2)
        self.assertFalse(eligible[1])

    def test_independent_can_win_only_one_fixed_seat(self):
        allocation, _ = danish.fixed_seats(np.array([1000, 100]), 3,
            np.array([True, False]), dhondt_gen, self.rng)
        np.testing.assert_array_equal(allocation, [1, 2])
        totals = danish.party_totals(np.array([1000, 100]), allocation,
            np.array([False, True]), 5, hare, "Quota", self.rng)
        np.testing.assert_array_equal(totals, [1, 4])

    def test_overhang_and_original_entitlement_cap(self):
        # Initial totals [1,2,6,1,10]. Removing overhung A would give B
        # three seats; cap B at its original two and recalculate again.
        totals = danish.party_totals(np.array([7, 24, 53, 11, 93]),
            np.array([2, 0, 4, 0, 8]), np.ones(5, bool), 20, hare, "Quota", self.rng)
        np.testing.assert_array_equal(totals, [2, 2, 5, 1, 10])

    def test_official_2022_overhang(self):
        fixture = json.loads((DATA / "denmark/national-results_2022.json").read_text())
        totals = danish.party_totals(np.array(fixture["votes"]),
            np.array(fixture["fixed_seats"]), np.array(fixture["eligible"]),
            175, hare, "Quota", self.rng)
        np.testing.assert_array_equal(totals, fixture["total_seats"])

    def test_zero_vote_regional_dead_end_is_explicit(self):
        with self.assertRaisesRegex(ValueError, "advance-allocation"):
            danish.prepare_regions(np.array([[10, 0], [0, 10]]), np.zeros((2, 2), int),
                [2, 0], [{"abbreviation": "H", "num_adj_seats": 1},
                         {"abbreviation": "SS", "num_adj_seats": 1}],
                [np.array([0]), np.array([1])], sainte_lague_gen, self.rng)

    def test_ties_are_reproducible_lots(self):
        results = []
        for _ in range(2):
            allocation, demo = danish.prepare_regions(np.array([[10, 10]]),
                np.zeros((1, 2), int), [1, 1], [{"abbreviation": "H", "num_adj_seats": 2}],
                [np.array([0])], sainte_lague_gen, np.random.default_rng(42))
            self.assertTrue(demo["data"][0]["lot"])
            results.append(demo["data"])
        self.assertEqual(*results)

    def test_simulation_excludes_independents_and_preserves_threshold_votes(self):
        original = deepcopy(self.table)
        prepared = simulation_vote_table(self.table, True)
        self.assertEqual(len(prepared["parties"]), 12)
        self.assertEqual(sum(prepared["pruned"]), 2436)
        self.assertEqual(self.table, original)
        settings = SimulationSettings()
        settings.update(simulation_count=3, cpu_count=1)
        simulation = Simulation(settings, [system_for(self.table)], self.table)
        simulation.simulate()
        self.assertEqual(simulation.iteration, 3)
        self.assertEqual(simulation.nparty, 12)
        self.assertEqual(sum(simulation.election_handler.elections[0].results["all_const_total"]), 175)

        with patch("simulate.generate_corr_votes", return_value=(np.zeros((10, 12)), [])):
            generated, _ = next(simulation.gen_votes())
            self.assertTrue(np.all(np.asarray(generated) >= 1))
        self.assertEqual(self.table, original)

    def test_simulation_drops_independents_in_comparisons_and_exports(self):
        from simulate import Sim_result
        from excel_util import simulation_to_xlsx
        systems = [system_for(self.table), system_for(self.table)]
        systems[0]["name"], systems[1]["name"] = "Denmark 1", "Denmark 2"
        systems[0]["compare_with"] = True
        settings = SimulationSettings()
        settings.update(simulation_count=2, cpu_count=1)
        simulation = Simulation(settings, systems, self.table)
        simulation.simulate()
        result = Sim_result(simulation.attributes())
        result.analysis()
        web_result = result.get_result_web(False)
        self.assertEqual(len(web_result["parties"]), 12)
        populated_groups = [
            group for group in web_result["vuedata"]["group_ids"]
            if web_result["vuedata"][group]
        ]
        displayed_measure = web_result["vuedata"][populated_groups[0]][0]["avg"][0]
        self.assertEqual(set(displayed_measure), {"value", "integer", "ci"})
        self.assertIsInstance(displayed_measure["value"], float)
        with TemporaryDirectory() as directory:
            path = Path(directory) / "simulation.xlsx"
            simulation_to_xlsx(web_result, path, {"fractional_digits": 2})
            book = load_workbook(path)
            self.assertIn("Party names", book.sheetnames)
            names = [row[1] for row in book["Party names"].iter_rows(values_only=True)]
            self.assertNotIn("Rashid Ali", names)
            self.assertTrue(any(
                cell.number_format == "#,##0.00"
                for sheet in book for row in sheet.iter_rows() for cell in row
            ))
            book.close()

    def test_unsupported_configuration_errors(self):
        system = system_for(self.table)
        system["adjustment_method"] = "switching"
        with self.assertRaisesRegex(ValueError, "requires Maximum"):
            ElectionHandler(self.table, [system], True)
        table = deepcopy(self.table)
        table.pop("regions")
        with self.assertRaisesRegex(ValueError, "requires a region table"):
            self.election(table)
